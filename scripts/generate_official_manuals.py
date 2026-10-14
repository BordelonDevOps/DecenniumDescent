#!/usr/bin/env python3
"""
Generate Decennium Descent official manuals (no foreign product branding):

  - Director's Manual (Dungeon Master Guide) → docs/manuals/
  - Bestiary of Aethel (Monster Manual, Volo-density target) → docs/manuals/

Sources: Player Rulebook extracts, cosmology/world lore, owner locks,
Bestiary.xlsx / bestiary_data_full.json, floor-encounters.json.

Banned in output prose: D&D, Dungeons & Dragons, Pathfinder, Volo, Trench Crusade,
Wizards, Paizo, SRD (as product), 5e/5th Edition as product names.
"""
from __future__ import annotations

import json
import re
from pathlib import Path

import openpyxl
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "manuals"
OUT.mkdir(parents=True, exist_ok=True)

BANNED = re.compile(
    r"\b(Dungeons?\s*&\s*Dragons|D&D|Pathfinder|Paizo|Wizards of the Coast|"
    r"Trench Crusade|Volo'?s?|5e|5th Edition|SRD|OGL)\b",
    re.I,
)

GATE_GUARDIANS = {
    "Vaskor",
    "The Bloom That Hungers",
    "Ashthirst the Hollow",
    "Gor Vekk",
    "Choralith",
    "Lumen of the Shattered Thread",
    "Xarog",
    "Eirfax",
    "Threxus",
    "The Marquis of Silence",
}


def scrub(text: str) -> str:
    if not text:
        return ""
    return BANNED.sub("[redacted]", str(text))


def add_heading(doc: Document, text: str, level: int = 1):
    doc.add_heading(scrub(text), level=level)


def add_para(doc: Document, text: str, *, bold: bool = False):
    p = doc.add_paragraph()
    run = p.add_run(scrub(text))
    run.bold = bold
    run.font.size = Pt(11)
    return p


def add_bullets(doc: Document, items: list[str]):
    for it in items:
        doc.add_paragraph(scrub(it), style="List Bullet")


def set_narrow_margins(doc: Document):
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)


def load_monsters() -> list[dict]:
    jp = ROOT / "src" / "bestiary_data_full.json"
    data = json.loads(jp.read_text(encoding="utf-8"))
    if isinstance(data, list):
        monsters = data
    else:
        monsters = data.get("monsters") or data.get("entries") or []

    by_name: dict[str, dict] = {}
    for m in monsters:
        name = (m.get("name") or "").strip()
        if name:
            by_name[name] = m

    # Enrich from Excel when combat fields present
    xlsx = ROOT / "Bestiary.xlsx"
    if xlsx.exists():
        wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
        ws = wb["Monsters"]
        rows = ws.iter_rows(values_only=True)
        header = [str(h) if h else f"c{i}" for i, h in enumerate(next(rows))]
        idx = {h: i for i, h in enumerate(header)}

        def cell(r, key, default=None):
            i = idx.get(key)
            if i is None or i >= len(r):
                return default
            return r[i] if r[i] is not None else default

        for r in rows:
            name = cell(r, "Name")
            if not name:
                continue
            name = str(name).strip()
            m = by_name.get(name)
            if not m:
                m = {"name": name}
                by_name[name] = m
            m["size"] = cell(r, "Size") or m.get("size")
            m["type"] = cell(r, "Type") or m.get("type")
            m["alignment"] = cell(r, "Alignment") or m.get("alignment")
            m["challengeRating"] = cell(r, "Challenge") or m.get("challengeRating")
            m["xp"] = cell(r, "XP") or m.get("xp")
            m["zoneFound"] = cell(r, "Zone Found") or m.get("zoneFound")
            m["armorClass"] = cell(r, "AC") or m.get("armorClass")
            m["hitPoints"] = cell(r, "HP") or m.get("hitPoints")
            m["speed"] = cell(r, "Speed") or m.get("speed")
            m["excelTraits"] = cell(r, "Traits")
            m["combatProvenance"] = cell(r, "Combat Provenance")
            m["needsReview"] = cell(r, "NeedsReview")
            m["attack1"] = {
                "name": cell(r, "Attack1 Name"),
                "bonus": cell(r, "Attack1 Bonus"),
                "damage": cell(r, "Attack1 Damage"),
                "type": cell(r, "Attack1 Type"),
            }
            m["attack2"] = {
                "name": cell(r, "Attack2 Name"),
                "bonus": cell(r, "Attack2 Bonus"),
                "damage": cell(r, "Attack2 Damage"),
                "type": cell(r, "Attack2 Type"),
            }
            m["multiattack"] = cell(r, "Multiattack")
            m["saveDC"] = cell(r, "Save DC")
        wb.close()
    monsters = list(by_name.values())
    monsters.sort(key=lambda m: (m.get("name") or "").lower())
    return monsters


def load_json(path: Path, default=None):
    if not path.exists():
        return default
    return json.loads(path.read_text(encoding="utf-8"))


def read_text_guess(path: Path, max_chars: int | None = None) -> str:
    if not path.exists():
        return ""
    raw = path.read_bytes()
    text = None
    # UTF-16 BOM
    if raw.startswith(b"\xff\xfe") or raw.startswith(b"\xfe\xff"):
        for enc in ("utf-16", "utf-16-le", "utf-16-be"):
            try:
                text = raw.decode(enc)
                break
            except UnicodeDecodeError:
                continue
    # UTF-16 without BOM (nulls on odd/even bytes)
    if text is None and len(raw) > 100:
        sample = raw[:4000]
        nulls = sample.count(0)
        if nulls > len(sample) * 0.3:
            for enc in ("utf-16-le", "utf-16-be", "utf-16"):
                try:
                    text = raw.decode(enc)
                    break
                except UnicodeDecodeError:
                    continue
    if text is None:
        for enc in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
            try:
                text = raw.decode(enc)
                break
            except UnicodeDecodeError:
                continue
    if text is None:
        text = raw.decode("utf-8", errors="replace")
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    if text.startswith("\ufeff"):
        text = text.lstrip("\ufeff")
    # Drop nulls if a mistaken decode left them
    if "\x00" in text[:200]:
        text = text.replace("\x00", "")
    if max_chars:
        return text[:max_chars]
    return text


def add_long_text(doc: Document, md_parts: list[str], text: str, *, chunk: int = 1800):
    """Add large extracts as sequential paragraphs for print bulk."""
    text = scrub(text).strip()
    if not text:
        add_para(doc, "(Extract empty or unavailable.)")
        md_parts.append("(Extract empty or unavailable.)\n")
        return
    # Prefer blank-line paragraphs; fall back to hard chunks
    blocks = [b.strip() for b in re.split(r"\n\s*\n", text) if b.strip()]
    if len(blocks) < 5:
        blocks = [text[i : i + chunk] for i in range(0, len(text), chunk)]
    for block in blocks:
        if len(block) < 20:
            continue
        # keep paragraphs printable
        for i in range(0, len(block), chunk):
            piece = block[i : i + chunk].strip()
            if piece:
                add_para(doc, piece)
                md_parts.append(piece + "\n\n")


# ---------------------------------------------------------------------------
# Director's Manual
# ---------------------------------------------------------------------------

def build_directors_manual(monsters: list[dict]) -> tuple[Document, str]:
    doc = Document()
    set_narrow_margins(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("DECENNIUM DESCENT")
    r.bold = True
    r.font.size = Pt(28)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("Director's Manual")
    r.bold = True
    r.font.size = Pt(22)

    tag = doc.add_paragraph()
    tag.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tag.add_run(
        "The official guide for Directors running expeditions into the dungeon beneath Aethel.\n"
        "Aligned to the Player Rulebook. Authoritative for table adjudication and simulation agents.\n"
        "Edition stamp: 2026-07-30"
    ).font.size = Pt(10)

    md_parts: list[str] = [
        "# Decennium Descent — Director's Manual\n",
        "*Official GM/Director reference. Aligned to the Player Rulebook. 2026-07-30.*\n",
        "*Simulation agents: treat this book as primary adjudication reference alongside the Player Rulebook.*\n",
    ]

    def chapter(title_text: str, body_paras: list[str], bullets: list[str] | None = None):
        add_heading(doc, title_text, 1)
        md_parts.append(f"\n## {title_text}\n")
        for p in body_paras:
            add_para(doc, p)
            md_parts.append(p + "\n")
        if bullets:
            add_bullets(doc, bullets)
            for b in bullets:
                md_parts.append(f"- {b}\n")

    # --- Ch 1
    chapter(
        "Chapter 1 — The Director's Charge",
        [
            "You are the Director. You adjudicate the Weave, the dungeon clock, and the consequences of mortal ambition. "
            "Players run characters; you run the world that answers them.",
            "Decennium Descent is a tabletop roleplaying companion set on Aethel during the Decennial Opening—the hundred-day window "
            "when the dungeon accepts expeditions. Your job is not to defeat the players. Your job is to make every choice matter: "
            "combat, research, craft, betrayal, mercy, and retreat.",
            "Tone: grim liturgy and high stakes. Heaven and Hell hold a deliberate tension. Mortals dig for glory and gold. "
            "Beneath both stories, older engines keep score. Present the public myth as complete at campaign start. "
            "Reveal deeper machinery only through contradictions the party earns.",
            "This manual never replaces the Player Rulebook for player-facing procedures. When they conflict, stop and ask the product owner "
            "before overwriting published rulebook text. Companion apps and playtest sims are mirrors and tools—not higher law.",
        ],
        [
            "Primary player text: decennium descent player rulebook.docx",
            "Primary Director text: this manual (plus Matrix of Oblivion, Cosmology, World of Aethel, Floor Encounters)",
            "Simulation agents must prefer this manual + Player Rulebook over ad-hoc playtest reports",
        ],
    )

    # --- Ch 2
    chapter(
        "Chapter 2 — Session Architecture",
        [
            "A session is a slice of expedition time. Track three clocks whenever relevant: calendar day within the Opening, "
            "hours toward the horn milestone, and fictional morning/afternoon/evening periods for team assignments.",
            "Open with stakes (why this floor, why today). Close with consequences (loot, injuries, maps, reputation, horn hours spent). "
            "Between, rotate spotlight across Combat, Recon, Research, Resource, and Support so non-fighters still matter.",
            "Use Quick Resolution for routine forage and maintenance, Scene-Based play for pivotal discoveries, Parallel Play when teams split, "
            "and Between-Session resolution for long research or craft chains.",
            "Paper tracks that the companion Session tool may not yet enforce: death saves, Hit Dice on short rest, concentration, "
            "expedition team assignment, and period planners. Adjudicate from the Player Rulebook and this manual until the app catches up.",
        ],
        [
            "Morning / Afternoon / Evening = one expedition day",
            "Long rest underground advances the Opening clock (standard table assumption: +8 hours unless a haven rule says otherwise)",
            "Label draft combat stats and draft XP awards as provisional when using companion sandbox numbers",
        ],
    )

    # --- Ch 3 Opening & Horn
    chapter(
        "Chapter 3 — The Opening, the Horn, and Point of No Return",
        [
            "Canon Opening Window: one hundred days — twenty-four hundred hours — once the Shardspire activates. "
            "When the window ends, the dungeon seals. Survivors still inside may remain trapped until the next Cycle.",
            "Horn / Point of No Return: at one thousand two hundred hours — day fifty — the halfway mark of the Opening, "
            "the horn sounds. Treat this as the point of no return for the cycle's push: retreat with what you have, "
            "or accept the sealed-depths risk toward the Cycle's end.",
            "Older companion code may still display a one-hundred-twenty-hour underground horn. That figure is superseded for Director canon "
            "unless the owner later restores it as a nested micro-budget. Prefer 2400 / 1200 in new simulations.",
            "Never merge Opening length with a single delve's combat pacing. Expeditions can leave and re-enter during the Opening "
            "unless a haven, seal, or story event forbids it. The Opening is the season; the delve is the week.",
        ],
        [
            "Opening = 100 days / 2400 hours",
            "Horn PoNR = day 50 / 1200 hours",
            "Seal at Opening end; next chance in ten years",
        ],
    )

    # --- Ch 4 Expeditions
    chapter(
        "Chapter 4 — Expeditions (Player Rulebook RAW)",
        [
            "Adventuring companies are Expeditions. Maximum roster: twelve characters working toward common goals—not merely who stands in initiative.",
            "Assign members across five teams. Optimal bands: Combat 4–6; Reconnaissance 1–3; Research 1–3; Resource 2–4; Support 1–3. "
            "Characters may shift teams by period. Multi-period tasks (deep research, long scouting) consume consecutive blocks.",
            "Activity results: Combat yields advancement, treasure, reputation, and monster parts. Recon yields maps and advantage. "
            "Research yields knowledge and identification. Resource yields food, materials, and craft components. "
            "Support extends duration, maintains gear, and recovers the wounded.",
            "Forces larger than twelve are multiple coordinating Expeditions, not one illegal roster. Coalition play is allowed; "
            "each Expedition still respects the twelve-seat organizational cap.",
            "Base camps: Temporary Camp, Fortified Outpost, Research Station, Forward Base. Upgrade with reputation, tools, and NPC specialists.",
        ],
    )

    # --- Ch 5 World
    chapter(
        "Chapter 5 — Aethel and the Dungeon",
        [
            "Aethel is the mortal world wrapped around a prison-engine. Public doctrine teaches that Calan'thir and Asmodeus jointly built "
            "the Crucible—also called Calan'thir's Trial—as neutral ground, proving ground, and monument of cooperation when existence itself was threatened.",
            "The dungeon is one hundred floors in ten Gate divisions. Each tier tests a principle of living difference. Gate Guardians stand at the seams. "
            "Use live bestiary names for those Guardians: Vaskor; The Bloom That Hungers; Ashthirst the Hollow; Gor Vekk; Choralith; "
            "Lumen of the Shattered Thread; Xarog; Eirfax; Threxus; The Marquis of Silence.",
            "Civilizations and cultures occupy upper and middle bands. They are not accidents—they generate traffic, conflict, and return traffic to the surface.",
            "Umbrin peoples and geography use locked product names only. Do not import foreign elven underworld branding. "
            "Sunlight Sensitivity and reputation rules remain as in the Player Rulebook and Umbrin lore stamps.",
        ],
    )

    # Load floor summary
    floors = load_json(ROOT / "src" / "data" / "floor-encounters.json", {})
    add_heading(doc, "Tier Gazetteer (Director Brief)", 2)
    md_parts.append("\n### Tier Gazetteer (Director Brief)\n")
    tiers = floors.get("tiers") or floors.get("zones") or []
    if isinstance(floors, dict) and not tiers:
        # try common shapes
        for k, v in floors.items():
            if isinstance(v, list) and v and isinstance(v[0], dict) and "name" in v[0]:
                tiers = v
                break
    if isinstance(tiers, list):
        for t in tiers[:12]:
            name = t.get("name") or t.get("tierName") or "Unknown Tier"
            add_heading(doc, str(name), 3)
            md_parts.append(f"\n#### {name}\n")
            desc = t.get("description") or t.get("summary") or ""
            if desc:
                add_para(doc, desc)
                md_parts.append(desc + "\n")
            for poi in (t.get("pointsOfInterest") or t.get("pois") or [])[:8]:
                pname = poi.get("name", "POI")
                pdesc = poi.get("description") or ""
                line = f"{pname}: {pdesc}" if pdesc else pname
                doc.add_paragraph(scrub(line), style="List Bullet")
                md_parts.append(f"- {line}\n")

    # --- Ch 6 Cosmology
    chapter(
        "Chapter 6 — Cosmology for Directors (Public and Secret)",
        [
            "Public complete explanation at campaign start: after the Sundering, Calan'thir and Asmodeus forged the Divine Pact and created the Crucible. "
            "Cultures argue which sovereign deserves more credit. Virtually everyone accepts joint creation.",
            "Director-secret stratum: the Silent Ones are the Old Ones. They built the Heartcell—the hundred-floor resonance machine, Gate divisions, "
            "opening cycle, and reward architecture—with loot as bait so mortals would voluntarily return, discover, and eventually release them. "
            "Calan'thir and Asmodeus discovered, claimed, modified, activated, or bound themselves to that machine; they did not invent it.",
            "Formulation to remember: The Silent Ones built the prison. Calan'thir and Asmodeus taught the world to call it a dungeon.",
            "The Divine Pact (the Pact) keeps them hidden by sustaining dissonance—war, strife, ambition. Absolute peace risks awakening the Old Ones "
            "to silence the universe again. The Matrix of Oblivion describes the reverse chain that collapses difference into stillness.",
            "Reveal secrets through floors: architecture older than script; treasure engineered to shape surface behavior; records calling the sovereigns "
            "Custodians, Bearers, or Signatories; evidence the Pact governs but did not originate the machine; Heartcell operating truths near the core.",
            "Mortal faiths (Fifteenfold Pantheon, planar diagrams taught in academies) are folk religion—useful for Clerics and culture, not objective cosmology. "
            "True hierarchy for Directors: Spiral; Old Ones / Ul'nak; Heartcell; Calan'thir and Asmodeus as custodians; First Choir (seven); "
            "Fifteenfold Pantheon as masks; cultural gods and ascended beings.",
        ],
    )

    # --- Ch 7 Resonance
    chapter(
        "Chapter 7 — Resonance, Oracle, and Magic Adjudication",
        [
            "Resonance Points (RP) are the Weave caster economy. They are not spell slots. Teach this in session zero. "
            "Cost scales with effect tier, affinity, and reputation modifiers per Resonance canon. Retune recovers RP on a short rest; "
            "full harmony restores on a long rest. Overcast and Burnout exist—label Burnout rules as draft if your table uses the stub.",
            "Oracle never spends RP. Oracle spends Revelation Charges (RC) on a separate track. Do not collapse the languages.",
            "Classes without an RP pool (including Oracle, Soldier, and Warrior as listed in Resonance materials) must not be forced into RP bookkeeping.",
            "Alignment affinity on RP spends matters when sheets record alignment. Log it. Novices will call RP 'slots'—correct gently and move on.",
        ],
    )

    # --- Ch 8 Combat
    chapter(
        "Chapter 8 — Combat, Threat, and Gate Guardians",
        [
            "Use the Bestiary of Aethel for creature statistics. Many companion combat blocks are draft sandbox numbers—label them at the table. "
            "Gate Guardians use live bestiary names and should feel like tier seals, not random bosses.",
            "Estimate pressure with expected-value thinking: rounds ≈ foe hit points ÷ party damage per round; incoming damage across those rounds "
            "compared to party hit points. Adjust for terrain, action denial, and support teams waiting at camp.",
            "Mass combat with full Expeditions: keep at most about eight creatures acting in detailed initiative at once; abstract the rest as "
            "reinforcement clocks, morale, or off-map objectives.",
            "Death, dying, and rests follow the Player Rulebook. If the companion Session tool lacks a track, use paper.",
        ],
        [f"Canonical Gate Guardians: {', '.join(sorted(GATE_GUARDIANS))}"],
    )

    # --- Ch 9 Treasure
    chapter(
        "Chapter 9 — Treasure, Craft, and Research",
        [
            "Loot is bait and fuel. Removing dungeon-touched objects seeds future wars, bloodlines, and return delvers. "
            "Award treasure that creates next session's problems.",
            "Tradeskills and crafting use the companion recipe corpus. Empty component lists in data are a known gap—assign salvage tables "
            "at the table rather than inventing silent infinite free crafts.",
            "Research teams translate, identify, and map. Spire lore, gate runes, and Surveyor seals are expedition currency. "
            "Between-session research is encouraged.",
        ],
    )

    # --- Ch 10 NPCs
    chapter(
        "Chapter 10 — Factions, Reputation, and NPCs",
        [
            "Reputation tracks open doors and close others. Academic factions care about research finds. Temple and infernal contacts "
            "care about Pact politics. Surface empires care about who returns with weapons that rewrite local power.",
            "Recruit NPC specialists into Support and Resource seats when the roster has room under twelve. "
            "Named allies should want something the dungeon can break.",
        ],
    )

    # --- Ch 11 Hazards
    chapter(
        "Chapter 11 — Hazards, Camps, and Logistics",
        [
            "Watch rotations fail; sun-shields deplete in radiant tiers; spore and heat zones punish unprepared Resource teams. "
            "Forward camps are operational centers: rest, storage, workshop, and rally point.",
            "Track rations, tool quality, and injury triage. Choir Covenant-style support authority (healers deciding who stands) "
            "prevents argument spirals at low hit points.",
        ],
    )

    # --- Ch 12 Creating content
    chapter(
        "Chapter 12 — Creating Floors, Foes, and Hooks",
        [
            "New content should serve the machine's appetite for difference: conflict, choice, and change. "
            "Avoid perfect peace as a campaign end-state unless you intend apocalypse.",
            "When inventing monsters, add them to Bestiary.xlsx and re-ingest. Prefer Excel as the portable source of truth for statistics.",
            "Hooks that only reward murder are thinner than hooks that force research, craft, or political fallout on the surface.",
        ],
    )

    # --- Ch 13 Simulation agents
    chapter(
        "Chapter 13 — Protocol for Simulation Agents",
        [
            "Future automated or assisted simulations must treat this Director's Manual and the Player Rulebook as primary references. "
            "Playtest reports are historical evidence, not overrides.",
            "Always: respect max twelve per Expedition (or explicitly label multi-expedition coalitions); use 2400/1200 Opening clocks; "
            "label draft combat/XP; keep Silent Ones/Heartcell out of player-facing summaries; use bestiary Gate Guardian names; "
            "exercise Combat/Recon/Research/Resource/Support when claiming a 'full expedition' sim.",
            "Never: invent foreign product cosmology; silently overwrite rulebook text; treat Appendix B folk pantheon as objective cosmic truth; "
            "call Resonance 'spell slots' in final reports without noting the correction.",
            "On conflict with published Player Rulebook or this manual: stop and ask the owner with (a) what the book says, "
            "(b) what would change, (c) options.",
        ],
    )

    # --- Appendix quick ref
    add_heading(doc, "Appendix A — Quick Reference Tables", 1)
    md_parts.append("\n## Appendix A — Quick Reference Tables\n")
    add_para(doc, f"Bestiary creatures currently indexed: {len(monsters)}.")
    md_parts.append(f"Bestiary creatures currently indexed: {len(monsters)}.\n")
    add_para(
        doc,
        "Expedition teams: Combat 4–6 · Recon 1–3 · Research 1–3 · Resource 2–4 · Support 1–3 · Cap 12.",
    )
    add_para(doc, "Opening 2400 h · Horn PoNR 1200 h · Cycle 10 years.")
    add_para(
        doc,
        "Source index: Player Rulebook; Matrix of Oblivion; Cosmology of Decennium Descent; World of Aethel; "
        "Floor Encounters; Bestiary.xlsx; docs/mechanics/* locks.",
    )

    # Embed substantial lore extracts to approach reference page-counts
    add_heading(doc, "Appendix B — Cosmology Extract (Director)", 1)
    md_parts.append("\n## Appendix B — Cosmology Extract (Director)\n")
    add_para(
        doc,
        "Full cosmology novel extract follows for Director study. Folk pantheon material remains player-facing folk religion; "
        "use Chapter 6 for secret hierarchy.",
    )
    add_long_text(doc, md_parts, read_text_guess(ROOT / "The_Cosmology_of_Decennium_Descent_EXTRACT.txt"))

    add_heading(doc, "Appendix C — World of Aethel Extract (Director)", 1)
    md_parts.append("\n## Appendix C — World of Aethel Extract (Director)\n")
    add_long_text(doc, md_parts, read_text_guess(ROOT / "The_World_of_Aethel_Lore_EXTRACT.txt"))

    add_heading(doc, "Appendix D — Matrix of Oblivion (Director)", 1)
    md_parts.append("\n## Appendix D — Matrix of Oblivion (Director)\n")
    add_long_text(doc, md_parts, read_text_guess(ROOT / "docs" / "playtest" / "MATRIX_OF_OBLIVION_EXTRACT.md"))

    add_heading(doc, "Appendix E — Silent Ones / Heartcell (Director-Secret)", 1)
    md_parts.append("\n## Appendix E — Silent Ones / Heartcell (Director-Secret)\n")
    add_long_text(doc, md_parts, read_text_guess(ROOT / "docs" / "mechanics" / "SILENT_ONES_CANON.md"))
    add_long_text(doc, md_parts, read_text_guess(ROOT / "docs" / "mechanics" / "OWNER_COSMOLOGY_LOCKS_2026-07-30.md"))

    add_heading(doc, "Appendix F — Descent Window & Expedition Canon", 1)
    md_parts.append("\n## Appendix F — Descent Window & Expedition Canon\n")
    add_long_text(doc, md_parts, read_text_guess(ROOT / "docs" / "mechanics" / "DESCENT_WINDOW_CANON.md"))
    add_long_text(doc, md_parts, read_text_guess(ROOT / "docs" / "mechanics" / "EXPEDITION_STRUCTURE_CANON.md"))

    add_heading(doc, "Appendix G — Resonance Canon (Director)", 1)
    md_parts.append("\n## Appendix G — Resonance Canon (Director)\n")
    add_long_text(doc, md_parts, read_text_guess(ROOT / "docs" / "mechanics" / "RESONANCE_CANON.md"))

    add_heading(doc, "Appendix H — Player Rulebook: Expeditions Chapter (mirror)", 1)
    md_parts.append("\n## Appendix H — Player Rulebook: Expeditions Chapter (mirror)\n")
    add_para(
        doc,
        "Extracted from the Player Rulebook living text for Director convenience. "
        "If this mirror drifts, the Player Rulebook .docx wins until the owner authorizes an edit.",
    )
    rb = read_text_guess(ROOT / "rulebook_output.txt")
    # Pull Expeditions section if present
    m = re.search(
        r"Expeditions in Aethel[\s\S]{0,25000}?Conclusion[\s\S]{0,2000}?(?=\n[A-Z][^\n]{0,80}\n|$)",
        rb,
        re.I,
    )
    if m:
        add_long_text(doc, md_parts, m.group(0))
    else:
        # fallback: search line window
        lines = rb.split("\n")
        start = next((i for i, ln in enumerate(lines) if "Expeditions in Aethel" in ln), None)
        if start is not None:
            add_long_text(doc, md_parts, "\n".join(lines[start : start + 200]))
        else:
            add_para(doc, "Expeditions chapter not located in rulebook_output.txt — open the Player Rulebook directly.")

    add_heading(doc, "Appendix I — Alignment Checklist vs Player Rulebook", 1)
    md_parts.append("\n## Appendix I — Alignment Checklist vs Player Rulebook\n")
    checks = [
        "Expedition max 12 and five teams match Player Rulebook Expeditions in Aethel.",
        "Resonance taught as pool economy, not foreign slot language.",
        "Umbrin naming and sunlight rules match current Player Rulebook geography pass.",
        "Appendix B folk religion in Player Rulebook remains player-facing until owner approves rewrite; Directors use Chapter 6 hierarchy.",
        "Combat action economy follows Player Rulebook turn structure.",
        "Opening 2400h and Horn PoNR 1200h are Director canon (Session HORN_AT=1200).",
    ]
    add_bullets(doc, checks)
    for c in checks:
        md_parts.append(f"- {c}\n")

    add_heading(doc, "Appendix J — Floor Encounters Digest", 1)
    md_parts.append("\n## Appendix J — Floor Encounters Digest\n")
    fe = load_json(ROOT / "src" / "data" / "floor-encounters.json", {})
    for t in fe.get("tiers") or []:
        add_heading(doc, str(t.get("name") or "Tier"), 2)
        md_parts.append(f"\n### {t.get('name')}\n")
        if t.get("description"):
            add_para(doc, t["description"])
            md_parts.append(t["description"] + "\n")
        for poi in t.get("pointsOfInterest") or []:
            line = f"{poi.get('name')}: {poi.get('description') or ''}".strip()
            doc.add_paragraph(scrub(line), style="List Bullet")
            md_parts.append(f"- {line}\n")
        for enc in (t.get("encounters") or t.get("sampleEncounters") or [])[:12]:
            if isinstance(enc, dict):
                line = f"Encounter: {enc.get('name') or enc}"
            else:
                line = f"Encounter: {enc}"
            doc.add_paragraph(scrub(str(line)), style="List Bullet")
            md_parts.append(f"- {line}\n")

    add_heading(doc, "Appendix K — Player Rulebook Concordance Extract (Director)", 1)
    md_parts.append("\n## Appendix K — Player Rulebook Concordance Extract (Director)\n")
    add_para(
        doc,
        "Large living extract from the Player Rulebook for Director cross-check. "
        "This is a mirror for search and simulation agents — the .docx remains authoritative. "
        "Do not edit the Player Rulebook from this appendix without owner approval.",
    )
    rb_full = read_text_guess(ROOT / "rulebook_output.txt", max_chars=450_000)
    add_long_text(doc, md_parts, rb_full)

    add_heading(doc, "Appendix L — Folk Religion Brief (Appendix B rewrite intent)", 1)
    md_parts.append("\n## Appendix L — Folk Religion Brief\n")
    add_long_text(doc, md_parts, read_text_guess(ROOT / "docs" / "mechanics" / "APPENDIX_B_FOLK_RELIGION.md"))

    add_heading(doc, "Appendix M — World / Dungeon Loop", 1)
    md_parts.append("\n## Appendix M — World / Dungeon Loop\n")
    add_long_text(doc, md_parts, read_text_guess(ROOT / "docs" / "mechanics" / "WORLD_DUNGEON_LOOP.md"))

    add_heading(doc, "Appendix N — Newcomer Director Checklist", 1)
    md_parts.append("\n## Appendix N — Newcomer Director Checklist\n")
    add_long_text(doc, md_parts, read_text_guess(ROOT / "docs" / "playtest" / "NEWCOMER_GM_GUIDE.md"))

    return doc, "".join(md_parts)


# ---------------------------------------------------------------------------
# Monster Manual / Bestiary of Aethel
# ---------------------------------------------------------------------------

LORE_BLURBS = {
    "Aberration": "Born where the Spiral thins and difference curdles into wrong geometry. Directors should emphasize incomprehension and environmental wrongness.",
    "Beast": "Natural or dungeon-twisted animals. Use them as resource pressure and travel threats as often as set-piece fights.",
    "Celestial": "Agents or echoes of higher harmony. Rarely simple allies—Pact politics stain even radiant motives.",
    "Construct": "Forged purpose. In the Anvil tiers, constructs are labor, weapons, and evidence of older engineering.",
    "Dragon": "Apex predators of ambition. Their hoards seed surface wars for generations.",
    "Elemental": "Concentrated states of matter and energy along Spiral folds. Terrain is half the encounter.",
    "Fey": "Adjacent-fold beings. Bargains cost more than steel.",
    "Fiend": "Chaos-Realm lineages—devils of bargained discord and demons of uncontrolled fracture. Their eternal war feeds dissonance the Pact requires.",
    "Giant": "Scale and siege. Use them to break camps and force Expeditions to choose what to save.",
    "Humanoid": "Peoples of Aethel and the upper floors. Politics matter; slaughter has reputational cost.",
    "Monstrosity": "Twisted life that should not breed true—until it does.",
    "Ooze": "Digestive hazards. Punish poor watch rotations and sealed corridors.",
    "Plant": "Bloomdeep and verdant hells. Area control and attrition.",
    "Undead": "History that refuses silence. Necrotic pressure tests Support teams.",
}


def lore_for(monster: dict) -> str:
    name = monster.get("name") or "Unknown"
    typ = (monster.get("type") or "Monstrosity").title()
    zone = monster.get("zoneFound") or "the deeper dark"
    base = LORE_BLURBS.get(typ, "A creature of the Decennium's appetite for conflict and change.")
    cr = monster.get("challengeRating")
    align = monster.get("alignment") or "unknown alignment"
    desc = scrub(monster.get("description") or "")[:600]
    if name in GATE_GUARDIANS:
        return (
            f"{name} is a Gate Guardian—a seal-intelligence woven into the Heartcell's public face as Crucible myth. "
            f"Directors should stage this foe as a tier threshold with campaign consequences: failed gates stall Opening progress; "
            f"victories reactivate seal machinery and scatter loot-as-bait into Aethel. Zone association: {zone}. "
            f"Challenge context {cr}. {base} "
            f"When running {name}, telegraph the tier's principle before the fight and ensure Support/Research teams have off-map work "
            f"so the Combat team is not the only story in the room."
        )
    para1 = f"Encountered about {zone}. Typical disposition: {align}. Challenge {cr}. {base}"
    para2 = (
        f"Directors: decide whether this creature is a patrol, a nest guardian, a harvest node, or a political actor. "
        f"Resource teams want parts; Research teams want behavior notes; Recon wants routes that avoid fair fights."
    )
    para3 = f"Field note: {desc}" if desc else (
        f"If statistics are draft/sandbox, announce that before initiative and prefer Excel edits over silent on-table invention."
    )
    return f"{para1} {para2} {para3}"


def format_attack(a: dict | None) -> str:
    if not a or not a.get("name"):
        return ""
    parts = [str(a.get("name"))]
    if a.get("bonus") is not None:
        parts.append(f"+{a.get('bonus')} to hit" if str(a.get("bonus")).lstrip("-").isdigit() or isinstance(a.get("bonus"), (int, float)) else f"bonus {a.get('bonus')}")
    if a.get("damage"):
        parts.append(f"{a.get('damage')} {a.get('type') or 'damage'}".strip())
    return " — ".join(parts)


def build_monster_manual(monsters: list[dict]) -> tuple[Document, str]:
    doc = Document()
    set_narrow_margins(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("DECENNIUM DESCENT")
    r.bold = True
    r.font.size = Pt(28)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("Bestiary of Aethel")
    r.bold = True
    r.font.size = Pt(22)

    tag = doc.add_paragraph()
    tag.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tag.add_run(
        "A Director's encyclopedia of foes, ecology, and Gate Guardians.\n"
        f"{len(monsters)} creatures compiled from Bestiary.xlsx and companion ingest.\n"
        "Combat blocks marked draft remain provisional sandbox statistics.\n"
        "Edition stamp: 2026-07-30"
    ).font.size = Pt(10)

    md_parts: list[str] = [
        "# Decennium Descent — Bestiary of Aethel\n",
        f"*Director monster encyclopedia. {len(monsters)} entries. 2026-07-30.*\n",
    ]

    add_heading(doc, "How to Use This Bestiary", 1)
    how = [
        "Each entry provides ecology notes for Directors, then a combat line derived from the portable Excel bestiary and companion ingest.",
        "Where Combat Provenance or companion flags say draft/sandbox, announce that at the table. Excel remains the editable source; re-ingest after changes.",
        "Gate Guardians are called out explicitly. Prefer their live names; do not revive retired cosmology combat aliases.",
        "Drops and crafting materials live on related Excel sheets—summarize at the table from the workbook when harvest matters.",
        "This book is patterned as a deep monster lore-and-stat volume for Decennium Descent alone. It contains no foreign product branding.",
    ]
    for h in how:
        add_para(doc, h)
        md_parts.append(h + "\n")

    add_heading(doc, "Gate Guardians of the Ten Divisions", 1)
    md_parts.append("\n## Gate Guardians of the Ten Divisions\n")
    add_para(
        doc,
        "These ten names are locked for gate encounters. Stage them as seals of the machine's public Crucible face.",
    )
    for g in sorted(GATE_GUARDIANS):
        doc.add_paragraph(scrub(g), style="List Bullet")
        md_parts.append(f"- {g}\n")

    # Type essays (deep lore-and-ecology chapters)
    add_heading(doc, "Ecologies of Aethel", 1)
    md_parts.append("\n## Ecologies of Aethel\n")
    for typ, blurb in sorted(LORE_BLURBS.items()):
        add_heading(doc, f"On {typ}s", 2)
        add_para(doc, blurb)
        add_para(
            doc,
            f"Count in current index: {sum(1 for m in monsters if (m.get('type') or '').lower() == typ.lower())}.",
        )
        md_parts.append(f"\n### On {typ}s\n{blurb}\n")

    add_heading(doc, "Creature Lexicon (A–Z)", 1)
    md_parts.append("\n## Creature Lexicon (A–Z)\n")

    for m in monsters:
        name = m.get("name") or "Unknown"
        add_heading(doc, name, 2)
        md_parts.append(f"\n### {name}\n")

        meta = (
            f"{m.get('size') or '?'} {m.get('type') or '?'}, {m.get('alignment') or 'unaligned'}. "
            f"Challenge {m.get('challengeRating') if m.get('challengeRating') is not None else '?'}"
            f" ({m.get('xp') if m.get('xp') is not None else '?'} XP). "
            f"Zone: {m.get('zoneFound') or 'unknown'}."
        )
        add_para(doc, meta)
        md_parts.append(meta + "\n")

        lore = lore_for(m)
        add_para(doc, lore)
        md_parts.append(lore + "\n")

        ac = m.get("armorClass")
        hp = m.get("hitPoints")
        spd = m.get("speed")
        combat_line = f"Armor {ac if ac is not None else '—'}; Hit Points {hp if hp is not None else '—'}; Speed {spd if spd is not None else '—'}."
        add_para(doc, combat_line, bold=True)
        md_parts.append(f"**{combat_line}**\n")

        a1 = format_attack(m.get("attack1"))
        a2 = format_attack(m.get("attack2"))
        if a1:
            add_para(doc, "Attack I: " + a1)
            md_parts.append(f"- Attack I: {a1}\n")
        if a2:
            add_para(doc, "Attack II: " + a2)
            md_parts.append(f"- Attack II: {a2}\n")
        if m.get("multiattack"):
            add_para(doc, f"Multiattack: {m.get('multiattack')}")
            md_parts.append(f"- Multiattack: {m.get('multiattack')}\n")
        if m.get("saveDC"):
            add_para(doc, f"Save DC: {m.get('saveDC')}")
            md_parts.append(f"- Save DC: {m.get('saveDC')}\n")

        traits = m.get("excelTraits") or ""
        if traits:
            add_para(doc, f"Traits: {traits}")
            md_parts.append(f"- Traits: {traits}\n")

        combat = m.get("combat") or {}
        if combat.get("draft"):
            add_para(doc, "Director note: companion combat block is draft/sandbox — verify before live table.")
            md_parts.append("*Draft/sandbox combat block.*\n")
        prov = m.get("combatProvenance")
        if prov:
            add_para(doc, f"Provenance: {prov}")
            md_parts.append(f"- Provenance: {prov}\n")

        if name in GATE_GUARDIANS:
            add_para(doc, "GATE GUARDIAN — treat as tier seal encounter.", bold=True)
            md_parts.append("**GATE GUARDIAN**\n")

    add_heading(doc, "Appendix — Index by Challenge", 1)
    md_parts.append("\n## Appendix — Index by Challenge\n")
    by_cr: dict[str, list[str]] = {}
    for m in monsters:
        cr = str(m.get("challengeRating") if m.get("challengeRating") is not None else "?")
        by_cr.setdefault(cr, []).append(m.get("name") or "?")
    for cr in sorted(by_cr.keys(), key=lambda x: (x == "?", x)):
        add_heading(doc, f"Challenge {cr}", 2)
        names = ", ".join(by_cr[cr][:80])
        if len(by_cr[cr]) > 80:
            names += f", … (+{len(by_cr[cr]) - 80} more)"
        add_para(doc, names)
        md_parts.append(f"\n### Challenge {cr}\n{names}\n")

    return doc, "".join(md_parts)


def estimate_pages(doc: Document) -> int:
    # rough: ~3 paragraphs per page at dense layout, or word count / 400
    words = 0
    for p in doc.paragraphs:
        words += len(p.text.split())
    return max(1, words // 400)


def main():
    print("Loading monsters…")
    monsters = load_monsters()
    print(f"  {len(monsters)} monsters")

    print("Building Director's Manual…")
    dm_doc, dm_md = build_directors_manual(monsters)
    dm_docx = OUT / "Decennium_Descent_Directors_Manual.docx"
    dm_md_path = OUT / "Decennium_Descent_Directors_Manual.md"
    dm_doc.save(dm_docx)
    dm_md_path.write_text(dm_md, encoding="utf-8")
    print(f"  wrote {dm_docx} (~{estimate_pages(dm_doc)} est. pages)")

    print("Building Bestiary of Aethel…")
    mm_doc, mm_md = build_monster_manual(monsters)
    mm_docx = OUT / "Decennium_Descent_Bestiary_of_Aethel.docx"
    mm_md_path = OUT / "Decennium_Descent_Bestiary_of_Aethel.md"
    mm_doc.save(mm_docx)
    mm_md_path.write_text(mm_md, encoding="utf-8")
    print(f"  wrote {mm_docx} (~{estimate_pages(mm_doc)} est. pages)")

    # Agent pointer
    pointer = OUT / "README_MANUALS.md"
    pointer.write_text(
        scrub(
            f"""# Official manuals (Director + Bestiary)

Generated 2026-07-30 from Player Rulebook-aligned canon, lore extracts, and `Bestiary.xlsx` / `bestiary_data_full.json`.

| Book | Files | Role |
|------|-------|------|
| **Director's Manual** | `Decennium_Descent_Directors_Manual.docx` / `.md` | Dungeon Master / Director reference for table + **simulation agents** |
| **Bestiary of Aethel** | `Decennium_Descent_Bestiary_of_Aethel.docx` / `.md` | Monster encyclopedia ({len(monsters)} entries) |

## Regeneration

```bash
python scripts/generate_official_manuals.py
```

## Continuity

- Do not put Silent Ones / Heartcell into player-facing books.
- Opening 2400 h · Horn PoNR 1200 h.
- Expedition max 12 (Player Rulebook RAW).
- No foreign tabletop-brand references in generated prose.
- Before overwriting Player Rulebook text, ask the owner (see `docs/mechanics/CONTINUITY_SOURCE_OF_TRUTH.md`).

Estimated pages are word-count heuristics (~400 words/page), not print-layout proofs.
"""
        ),
        encoding="utf-8",
    )
    print(f"  wrote {pointer}")


if __name__ == "__main__":
    main()
