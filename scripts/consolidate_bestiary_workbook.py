#!/usr/bin/env python3
"""
One-shot (or re-run) merge: combat columns from sandbox + metadata gaps from
final_bestiary.xlsx → canonical Bestiary.xlsx. Backs up before overwrite.
"""

from __future__ import annotations

import json
import shutil
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
PRIMARY_XLSX = ROOT / "Bestiary.xlsx"
SANDBOX_XLSX = ROOT / "Bestiary_combat_sandbox.xlsx"
METADATA_XLSX = ROOT / "final_bestiary.xlsx"
BACKUP_DIR = ROOT / "backups"

COMBAT_COLUMNS = [
    "AC",
    "HP",
    "Speed",
    "Attack1 Name",
    "Attack1 Bonus",
    "Attack1 Damage",
    "Attack1 Type",
    "Attack2 Name",
    "Attack2 Bonus",
    "Attack2 Damage",
    "Attack2 Type",
    "Multiattack",
    "Save DC",
    "Traits",
    "Combat Provenance",
]

METADATA_COLUMNS = [
    "Size",
    "Type",
    "Tags",
    "Alignment",
    "Challenge",
    "XP",
    "Zone Found",
    "Tier 1 Drop",
    "Tier 1 Rate",
    "Tier 1 Value",
    "Tier 2 Drop",
    "Tier 2 Rate",
    "Tier 2 Value",
    "Tier 3 Drop",
    "Tier 3 Rate",
    "Tier 3 Value",
    "Tier 4 Drop",
    "Tier 4 Rate",
    "Tier 4 Value",
    "Tier 5 Drop",
    "Tier 5 Rate",
    "Tier 5 Value",
    "Tier 6 Drop",
    "Tier 6 Rate",
    "Tier 6 Value",
]

import sys

sys.path.insert(0, str(ROOT / "scripts"))
from scrub_bestiary_ip import rename_value  # noqa: E402


def is_blank(val: Any) -> bool:
    if val is None:
        return True
    if isinstance(val, float) and pd.isna(val):
        return True
    text = str(val).strip()
    return not text or text.lower() == "nan"


def backup_primary() -> str:
    BACKUP_DIR.mkdir(exist_ok=True)
    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    dest = BACKUP_DIR / f"Bestiary_{stamp}.xlsx"
    shutil.copy2(PRIMARY_XLSX, dest)
    return str(dest)


def header_index(ws) -> dict[str, int]:
    out: dict[str, int] = {}
    for col in range(1, ws.max_column + 1):
        val = ws.cell(row=1, column=col).value
        if val:
            out[str(val).strip()] = col
    return out


def main() -> None:
    if not PRIMARY_XLSX.exists():
        raise SystemExit(f"Missing {PRIMARY_XLSX}")
    if not SANDBOX_XLSX.exists():
        raise SystemExit(f"Missing {SANDBOX_XLSX} — nothing to merge")

    backup = backup_primary()
    sandbox_df = pd.read_excel(SANDBOX_XLSX, sheet_name="Monsters")
    meta_df = (
        pd.read_excel(METADATA_XLSX, sheet_name="Monsters")
        if METADATA_XLSX.exists()
        else pd.DataFrame()
    )

    sandbox_by_name: dict[str, pd.Series] = {}
    for _, row in sandbox_df.iterrows():
        raw = row.get("Name")
        if not is_blank(raw):
            sandbox_by_name[rename_value(str(raw).strip())] = row

    meta_by_name: dict[str, pd.Series] = {}
    for _, row in meta_df.iterrows():
        raw = row.get("Name")
        if not is_blank(raw):
            meta_by_name[rename_value(str(raw).strip())] = row

    wb = load_workbook(PRIMARY_XLSX)
    ws = wb["Monsters"]
    headers = header_index(ws)

    # Ensure combat + any missing metadata headers exist
    next_col = ws.max_column + 1
    for col_name in METADATA_COLUMNS + COMBAT_COLUMNS:
        if col_name not in headers:
            ws.cell(row=1, column=next_col, value=col_name)
            headers[col_name] = next_col
            next_col += 1

    name_col = headers.get("Name", 1)
    name_to_row: dict[str, int] = {}
    for r in range(2, ws.max_row + 1):
        val = ws.cell(row=r, column=name_col).value
        if val:
            name_to_row[rename_value(str(val).strip())] = r

    meta_filled = 0
    combat_filled = 0
    for name, r in name_to_row.items():
        sandbox_row = sandbox_by_name.get(name)
        meta_row = meta_by_name.get(name)

        for col_name in METADATA_COLUMNS:
            col_idx = headers[col_name]
            current = ws.cell(row=r, column=col_idx).value
            if not is_blank(current):
                continue
            if meta_row is not None and col_name in meta_row.index and not is_blank(meta_row[col_name]):
                ws.cell(row=r, column=col_idx, value=meta_row[col_name])
                meta_filled += 1

        if sandbox_row is None:
            continue
        for col_name in COMBAT_COLUMNS:
            if col_name not in sandbox_row.index:
                continue
            val = sandbox_row[col_name]
            if is_blank(val):
                continue
            col_idx = headers[col_name]
            ws.cell(row=r, column=col_idx, value=val)
            combat_filled += 1

    wb.save(PRIMARY_XLSX)

    summary = {
        "canonical": str(PRIMARY_XLSX.name),
        "backup": backup,
        "rows": len(name_to_row),
        "metadata_cells_filled": meta_filled,
        "combat_cells_filled": combat_filled,
        "combat_columns": COMBAT_COLUMNS,
    }
    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main()
