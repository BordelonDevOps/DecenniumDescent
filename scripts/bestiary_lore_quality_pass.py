# -*- coding: utf-8 -*-
"""
Lore-aligned Bestiary.xlsx quality pass (owner sign-off 2026-07-30:
apply changes that follow established Decennium Descent lore).

- Gate Guardian Type/Zone/CR/Alignment/Description floors
- Resolve Xarog / guardian aliases (rename-only → canonical; delete if both)
- IP / deity-lock renames from Codex ledger
- Culture-family Type normalization
- Archetype Type fixes for clear name→type matches
- Type-flavored provisional attack names (replace generic sandbox verbs)
- Clear NeedsReview where we corrected the flagged issue

Does NOT invent cosmology. Backs up Bestiary.xlsx first.
"""
from __future__ import annotations

import re
import shutil
from datetime import datetime
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "Bestiary.xlsx"
BACKUP = ROOT / "backups" / f"Bestiary_pre_quality_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"

# --- Gate Guardians (Chronicle / Director locks) ---
# Seal ladder CR: monotonic Novice→Legendary (Director lore curve)
GUARDIANS = {
    "Vaskor, the Stone Maw": {
        "Type": "Elemental",
        "Size": "Huge",
        "Alignment": "Neutral",
        "Challenge": 8,
        "XP": 3900,
        "Zone Found": "Yieldheart Caverns",
        "Tags": "Gate Guardian",
        "floor": 10,
        "gate": "Warded Gate",
        "desc": "Colossal earth elemental seal of Possibility at the Warded Gate (Floor 10, Yieldheart Caverns).",
    },
    "The Bloom That Hungers": {
        "Type": "Plant",
        "Size": "Huge",
        "Alignment": "Chaotic Evil",
        "Challenge": 11,
        "XP": 7200,
        "Zone Found": "Bloomdeep Marsh",
        "Tags": "Gate Guardian",
        "floor": 20,
        "gate": "Sporemist Gate",
        "desc": "Spore-hunger seal of Choice at the Sporemist Gate (Floor 20, Bloomdeep Marsh).",
    },
    "Ashthirst the Hollow": {
        "Type": "Elemental",
        "Size": "Huge",
        "Alignment": "Neutral Evil",
        "Challenge": 14,
        "XP": 11500,
        "Zone Found": "Anvil of Brogi",
        "Tags": "Gate Guardian",
        "floor": 30,
        "gate": "Lavaflow Gate",
        "desc": "Hollow forge-elemental seal of Difference at the Lavaflow Gate (Floor 30, Anvil of Brogi).",
    },
    "Gor Vekk, Architect of Ends": {
        "Type": "Aberration",
        "Size": "Large",
        "Alignment": "Lawful Evil",
        "Challenge": 17,
        "XP": 18000,
        "Zone Found": "Isles of Eternal Radiance",
        "Tags": "Gate Guardian",
        "floor": 40,
        "gate": "Radiance Gate",
        "desc": "Ancient aberration seal of Conflict at the Radiance Gate (Floor 40, Isles of Eternal Radiance). Projection fights may use reduced HP — full CR 17 is the seal budget.",
    },
    "Gor Vekk": {  # alias if sheet uses short name
        "alias_of": "Gor Vekk, Architect of Ends"
    },
    "Choralith, Voice Below": {
        "Type": "Aberration",
        "Size": "Large",
        "Alignment": "Neutral",
        "Challenge": 20,
        "XP": 25000,
        "Zone Found": "Sanctum of Echoing Choirs",
        "Tags": "Gate Guardian",
        "floor": 50,
        "gate": "Canticle Gate",
        "desc": "Echo-seal of Motion at the Canticle Gate (Floor 50, Sanctum of Echoing Choirs). Soft quit line is morale geography — not the Opening horn PoNR.",
    },
    "Choralith": {"alias_of": "Choralith, Voice Below"},
    "Lumen of the Shattered Thread": {
        "Type": "Elemental",
        "Size": "Large",
        "Alignment": "Chaotic Neutral",
        "Challenge": 22,
        "XP": 41000,
        "Zone Found": "Crossroads of the Shattered Thread",
        "Tags": "Gate Guardian",
        "floor": 60,
        "gate": "Lumen Gate",
        "desc": "Shattered-Weave elemental seal of Time at the Lumen Gate (Floor 60, Crossroads of the Shattered Thread).",
    },
    "Xarog, the Pactbound Furnace": {
        "Type": "Fiend",
        "Size": "Huge",
        "Alignment": "Lawful Evil",
        "Challenge": 24,
        "XP": 62000,
        "Zone Found": "Pactbound Arena",
        "Tags": "Gate Guardian",
        "floor": 70,
        "gate": "Furnace Gate",
        "desc": "Pactbound fiend seal of History at the Furnace Gate (Floor 70, Pactbound Arena).",
    },
    # Sheet sometimes lacks the comma; resolve to canonical before GATE fix
    "Xarog the Pactbound Furnace": {"alias_of": "Xarog, the Pactbound Furnace"},
    "Eirfax, the Dusk Crowned": {
        "Type": "Fey",
        "Size": "Medium",
        "Alignment": "Neutral Evil",
        "Challenge": 26,
        "XP": 90000,
        "Zone Found": "Duskshard Pools",
        "Tags": "Gate Guardian",
        "floor": 80,
        "gate": "Crown Gate",
        "desc": "Twilight fey seal of Identity at the Crown Gate (Floor 80, Duskshard Pools).",
    },
    "Threxus, the Iron Spiral": {
        "Type": "Construct",
        "Size": "Gargantuan",
        "Alignment": "Lawful Neutral",
        "Challenge": 28,
        "XP": 120000,
        "Zone Found": "Veiled Iron Spiral",
        "Tags": "Gate Guardian",
        "floor": 90,
        "gate": "Spiral Gate",
        "desc": "Spiraling iron construct seal of Will at the Spiral Gate (Floor 90, Veiled Iron Spiral).",
    },
    "The Marquis of Silence": {
        "Type": "Fiend",
        "Size": "Huge",
        "Alignment": "Chaotic Evil",
        "Challenge": 30,
        "XP": 155000,
        "Zone Found": "Silencebound Void",
        "Tags": "Gate Guardian",
        "floor": 100,
        "gate": "Marquis Gate",
        "desc": "Fiend seal of Life's hush at the Marquis Gate (Floor 100, Silencebound Void). Rules through silence and fear.",
    },
}

# Exact renames (Codex ledger + deity locks) — owner lore-aligned sign-off
RENAMES = {
    "Flumph": "Lumen Drifter",
    "Owlbear": "Hookbeast",
    "Tarrasque": "Worldbreaker",
    "Intellect Devourer": "Mind Leech",
    "Neothelid": "Spiral Worm",
    "Draegloth": "Umbrin Warborn",
    "Drider": "Umbrin Arachnid",
    "Displacer Beast": "Refractor Panther",
    "Hook Horror": "Rip-Auger",
    "Umber Hulk": "Deep Borer",
    "Rust Monster": "Corrosion Grazer",
    "Carrion Crawler": "Charnel Crawler",
    "Aarakocra": "Skyfolk of the High Vents",
    "Kenku": "Mimic-Crow",
    "Tortle": "Shellkin",
    "Tortle Druid": "Shellkin Warden",
    "Thri-kreen": "Chitin-Strider",
    "Bullywug": "Bogkin",
    "Grung": "Poison-Skin",
    "Grung Elite Warrior": "Poison-Skin Elite",
    "Grung Wildling": "Poison-Skin Wildling",
    "Gnoll Fang of Yeenoghu": "Gnoll Fang of the Howling Fracture",
    "Orc Blade of Ilneval": "Orc Blade of the Iron March",
    "Orc Eye of Gruumsh": "Orc Eye of the One-Eyed Pact",
    "Orc Hand of Yurtrus": "Orc Hand of the Rot Sigil",
    "Orc Nurtured One of Yurtrus": "Orc Nurtured One of the Rot Sigil",
    "Orc Red Fang of Shargaas": "Orc Red Fang of the Veiled Hunt",
    "Orc Claw of Luthic": "Orc Claw of the Den-Mother",
    "Spawn of Kyuss": "Spawn of the Worm Spiral",
    "Firenewt Warlock of Imix": "Firenewt Warlock of the Ember Court",
    "Xvart Warlock of Raxivort": "Xvart Warlock of the Nine-Legged Choir",
    # Yugoloth → Hush-Broker
    "Merrenoloth": "Ferry-Wraith of the Hush",
    "Ultroloth": "Court Broker of the Hush",
    "Arcanaloth": "Hush-Broker Arcanaloth",
    "Canoloth": "Hush-Broker Canoloth",
    "Dhergoloth": "Hush-Broker Dhergoloth",
    "Hydroloth": "Hush-Broker Hydroloth",
    "Mezzoloth": "Hush-Broker Mezzoloth",
    "Nycaloth": "Hush-Broker Nycaloth",
    "Oinoloth": "Hush-Broker Oinoloth",
    "Yagnoloth": "Hush-Broker Yagnoloth",
    # Modron → Order-Drone
    "Duodrone": "Order-Drone Duodrone",
    "Monodrone": "Order-Drone Monodrone",
    "Pentadrone": "Order-Drone Pentadrone",
    "Quadrone": "Order-Drone Quadrone",
    "Tridrone": "Order-Drone Tridrone",
    "Marut": "Order-Drone Marut",
    # Abishai
    "Black Abishai": "Pact Devil (Black Court)",
    "Blue Abishai": "Pact Devil (Blue Court)",
    "Green Abishai": "Pact Devil (Green Court)",
    "Red Abishai": "Pact Devil (Red Court)",
    "White Abishai": "Pact Devil (White Court)",
    # Slaad
    "Blue Slaad": "Chaos-Husk Blue",
    "Death Slaad": "Chaos-Husk Death",
    "Gray Slaad": "Chaos-Husk Gray",
    "Green Slaad": "Chaos-Husk Green",
    "Red Slaad": "Chaos-Husk Red",
    "Slaad Tadpole": "Chaos-Husk Tadpole",
    # Eladrin
    "Autumn Eladrin": "Season-Bound Autumn",
    "Spring Eladrin": "Season-Bound Spring",
    "Summer Eladrin": "Season-Bound Summer",
    "Winter Eladrin": "Season-Bound Winter",
    # Sorrowsworn
    "Angry Sorrowsworn": "Angry Grief-Bound",
    "Hungry Sorrowsworn": "Hungry Grief-Bound",
    "Lonely Sorrowsworn": "Lonely Grief-Bound",
    "Lost Sorrowsworn": "Lost Grief-Bound",
    "Wretched Sorrowsworn": "Wretched Grief-Bound",
    # Star Spawn
    "Star Spawn Grue": "Spiral Breach Grue",
    "Star Spawn Hulk": "Spiral Breach Hulk",
    "Star Spawn Larva Mage": "Spiral Breach Larva Mage",
    "Star Spawn Mangler": "Spiral Breach Mangler",
    "Star Spawn Seer": "Spiral Breach Seer",
    # Neogi
    "Neogi": "Void Broker",
    "Neogi Hatchling": "Void Broker Hatchling",
    "Neogi Master": "Void Broker Master",
    # Demon hierarchy (Fracture-Kin)
    "Balor": "Fracture-Kin Balor",
    "Marilith": "Fracture-Kin Marilith",
    "Glabrezu": "Fracture-Kin Glabrezu",
    "Nalfeshnee": "Fracture-Kin Nalfeshnee",
    "Vrock": "Fracture-Kin Vrock",
    "Hezrou": "Fracture-Kin Hezrou",
    "Babau": "Fracture-Kin Babau",
    "Dretch": "Fracture-Kin Dretch",
    "Manes": "Fracture-Kin Manes",
    "Molydeus": "Fracture-Kin Molydeus",
    "Nabassu": "Fracture-Kin Nabassu",
    "Goristro": "Fracture-Kin Goristro",
    "Yochlol": "Fracture-Kin Yochlol",
    "Armanite": "Fracture-Kin Armanite",
    "Bulezau": "Fracture-Kin Bulezau",
    "Shoosuva": "Fracture-Kin Shoosuva",
    "Maurezhi": "Fracture-Kin Maurezhi",
    "Alkilith": "Fracture-Kin Alkilith",
    "Wastrilith": "Fracture-Kin Wastrilith",
    "Sibriex": "Fracture-Kin Sibriex",
    "Chasme": "Fracture-Kin Chasme",
    "Quasit": "Fracture-Kin Quasit",
    "Barlgura": "Fracture-Kin Barlgura",
    "Shadow Demon": "Fracture-Kin Shadow Demon",
    "Vargouille": "Fracture-Kin Vargouille",
}

# Prefix culture → Type
CULTURE_PREFIX_TYPE = [
    ("Umbrin", "Humanoid"),
    ("Deep Dwarf", "Humanoid"),
    ("Serpent-Blood", "Monstrosity"),
    ("Astral Reaver", "Humanoid"),
    ("Astral Ascetic", "Humanoid"),
    ("Depth Toad", "Humanoid"),
    ("Gaze Tyrant", "Aberration"),
    ("Voidmind", "Aberration"),
    ("Sporefolk", "Plant"),
    ("Cave Gnome", "Humanoid"),
    ("Poison-Skin", "Humanoid"),
    ("Shellkin", "Humanoid"),
    ("Order-Drone", "Construct"),
    ("Hush-Broker", "Fiend"),
    ("Ferry-Wraith", "Fiend"),
    ("Court Broker", "Fiend"),
    ("Fracture-Kin", "Fiend"),
    ("Pact Devil", "Fiend"),
    ("Chaos-Husk", "Aberration"),
    ("Season-Bound", "Fey"),
    ("Grief-Bound", "Monstrosity"),
    ("Spiral Breach", "Aberration"),
    ("Void Broker", "Aberration"),
]

# Exact name → Type for clear archetypes
EXACT_TYPE = {
    "Shadow": "Undead",
    "Bodak": "Undead",
    "Ghost": "Undead",
    "Ghoul": "Undead",
    "Specter": "Undead",
    "Wraith": "Undead",
    "Revenant": "Undead",
    "Vampire": "Undead",
    "Mummy": "Undead",
    "Zombie": "Undead",
    "Skeleton": "Undead",
    "Azer": "Elemental",
    "Quipper": "Beast",
    "Wolf": "Beast",
    "Boar": "Beast",
    "Lion": "Beast",
    "Tiger": "Beast",
    "Bear": "Beast",
    "Horse": "Beast",
    "Cow": "Beast",
    "Imp": "Fiend",
    "Pit Fiend": "Fiend",
    "Bone Devil": "Fiend",
    "Ice Devil": "Fiend",
    "Horned Devil": "Fiend",
    "Erinyes": "Fiend",
    "Aboleth": "Aberration",
    "Otyugh": "Aberration",
    "Cloaker": "Aberration",
    "Treant": "Plant",
    "Shambling Mound": "Plant",
    "Animated Armor": "Construct",
    "Clay Golem": "Construct",
    "Shield Guardian": "Construct",
    "Helmed Horror": "Construct",
    "Air Elemental": "Elemental",
    "Earth Elemental": "Elemental",
    "Fire Elemental": "Elemental",
    "Water Elemental": "Elemental",
    "Xorn": "Elemental",
    "Hookbeast": "Monstrosity",
    "Lumen Drifter": "Aberration",
    "Worldbreaker": "Monstrosity",
    "Mind Leech": "Aberration",
    "Spiral Worm": "Aberration",
    "Umbrin Warborn": "Fiend",
    "Umbrin Arachnid": "Monstrosity",
    "Refractor Panther": "Monstrosity",
    "Rip-Auger": "Monstrosity",
    "Deep Borer": "Monstrosity",
    "Corrosion Grazer": "Monstrosity",
    "Charnel Crawler": "Monstrosity",
    "Skyfolk of the High Vents": "Humanoid",
    "Mimic-Crow": "Humanoid",
    "Bogkin": "Humanoid",
    "Chitin-Strider": "Humanoid",
}

ATTACKS_BY_TYPE = {
    "Beast": ("Rending Bite", "Claw Swipe", "piercing", "slashing"),
    "Humanoid": ("War Blade", "Crossbow Bolt", "slashing", "piercing"),
    "Undead": ("Grave Touch", "Soul Chill", "necrotic", "cold"),
    "Fiend": ("Hellish Claw", "Infernal Lash", "slashing", "fire"),
    "Celestial": ("Radiant Smite", "Judgment Beam", "radiant", "radiant"),
    "Elemental": ("Elemental Slam", "Surge Burst", "bludgeoning", "force"),
    "Construct": ("Iron Fist", "Gear Shear", "bludgeoning", "slashing"),
    "Dragon": ("Drake Bite", "Tail Sweep", "piercing", "bludgeoning"),
    "Aberration": ("Warp Tentacle", "Mind Spike", "bludgeoning", "psychic"),
    "Monstrosity": ("Hooked Claw", "Crushing Maw", "slashing", "piercing"),
    "Ooze": ("Acid Pseudopod", "Engulf Pulse", "acid", "acid"),
    "Plant": ("Vine Lash", "Spore Burst", "bludgeoning", "poison"),
    "Fey": ("Glamour Blade", "Hex Dart", "slashing", "psychic"),
    "Giant": ("Tree-Trunk Club", "Boulder Hurl", "bludgeoning", "bludgeoning"),
}

GENERIC_ATTACKS = {"Freak Bite", "Monstrous Claw", "Aberrant Lash", "Tail Lash", "Grave Shriek"}


def colmap(ws):
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    return {str(h): i + 1 for i, h in enumerate(header) if h}


def ensure_desc_col(ws, cols):
    if "Description" in cols:
        return cols
    # append Description column
    new_i = len(cols) + 1
    ws.cell(1, new_i, "Description")
    cols["Description"] = new_i
    return cols


def main():
    BACKUP.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(XLSX, BACKUP)
    print("backup", BACKUP)

    wb = openpyxl.load_workbook(XLSX)
    ws = wb["Monsters"]
    cols = colmap(ws)
    cols = ensure_desc_col(ws, cols)

    def get(row, key, default=None):
        i = cols.get(key)
        if not i:
            return default
        return ws.cell(row, i).value

    def setc(row, key, value):
        i = cols.get(key)
        if not i:
            return
        ws.cell(row, i).value = value

    # Collect rows by name
    rows_by_name: dict[str, list[int]] = {}
    for r in range(2, ws.max_row + 1):
        name = get(r, "Name")
        if not name:
            continue
        rows_by_name.setdefault(str(name).strip(), []).append(r)

    log: list[str] = []
    deleted = 0

    def rebuild_name_map():
        nonlocal cols, rows_by_name
        cols = colmap(ws)
        cols = ensure_desc_col(ws, cols)
        rows_by_name = {}
        for r in range(2, ws.max_row + 1):
            name = get(r, "Name")
            if name:
                rows_by_name.setdefault(str(name).strip(), []).append(r)

    # Resolve guardian aliases (idempotent):
    # - alias only → rename to canonical, then GATE fix applies
    # - both alias + canonical → delete alias rows
    need_rebuild = False
    for gname, spec in GUARDIANS.items():
        if "alias_of" not in spec:
            continue
        canonical = spec["alias_of"]
        if gname not in rows_by_name:
            continue
        if canonical in rows_by_name:
            for r in sorted(rows_by_name[gname], reverse=True):
                ws.delete_rows(r, 1)
                deleted += 1
                log.append(f"DELETE duplicate alias: {gname} (kept {canonical})")
            need_rebuild = True
        else:
            for r in rows_by_name[gname]:
                setc(r, "Name", canonical)
                log.append(f"RENAME alias {gname} → {canonical}")
            rows_by_name[canonical] = rows_by_name.pop(gname)
    if need_rebuild:
        rebuild_name_map()

    # Apply renames (longest first)
    for old, new in sorted(RENAMES.items(), key=lambda x: -len(x[0])):
        if old not in rows_by_name:
            continue
        for r in rows_by_name[old]:
            setc(r, "Name", new)
            log.append(f"RENAME {old} → {new}")
        # update map
        rows_by_name[new] = rows_by_name.pop(old)

    # Rebuild name map after renames
    rows_by_name = {}
    for r in range(2, ws.max_row + 1):
        name = get(r, "Name")
        if name:
            rows_by_name.setdefault(str(name).strip(), []).append(r)

    # Gate guardians
    for gname, spec in GUARDIANS.items():
        if "alias_of" in spec:
            continue
        targets = rows_by_name.get(gname, [])
        # also try without comma variants already handled
        if not targets:
            log.append(f"WARN guardian missing: {gname}")
            continue
        for r in targets:
            for k in ("Type", "Size", "Alignment", "Challenge", "XP", "Zone Found", "Tags"):
                if k in spec:
                    setc(r, k, spec[k])
            if "Description" in cols:
                setc(r, "Description", spec["desc"])
            # bespoke attacks for guardians if generic
            a1 = get(r, "Attack1 Name")
            if not a1 or str(a1) in GENERIC_ATTACKS:
                typ = spec["Type"]
                a1n, a2n, t1, t2 = ATTACKS_BY_TYPE.get(
                    typ, ("Seal Strike", "Seal Pulse", "force", "force")
                )
                setc(r, "Attack1 Name", f"{gname.split(',')[0]}'s {a1n}")
                setc(r, "Attack1 Type", t1)
                if cols.get("Attack2 Name") and (
                    not get(r, "Attack2 Name") or str(get(r, "Attack2 Name")) in GENERIC_ATTACKS
                ):
                    setc(r, "Attack2 Name", a2n)
                    setc(r, "Attack2 Type", t2)
            setc(r, "NeedsReview", None)
            setc(r, "Combat Provenance", "lore-pass-2026-07-30-gate")
            # scale HP roughly with CR if absurdly low for new CR
            cr = spec["Challenge"]
            hp = get(r, "HP")
            try:
                hp_i = int(hp)
            except (TypeError, ValueError):
                hp_i = 0
            target_hp = max(hp_i, 20 + cr * 18)
            if hp_i < target_hp * 0.5:
                setc(r, "HP", target_hp)
            ac = get(r, "AC")
            try:
                ac_i = int(ac)
            except (TypeError, ValueError):
                ac_i = 10
            setc(r, "AC", max(ac_i, 12 + cr // 3))
            log.append(f"GATE fix {gname}: Type={spec['Type']} CR={cr} Zone={spec['Zone Found']}")

    # Culture + exact types + attacks for all rows
    for r in range(2, ws.max_row + 1):
        name = get(r, "Name")
        if not name:
            continue
        name = str(name).strip()
        typ = get(r, "Type")
        typ_s = str(typ).strip() if typ else ""

        # culture prefix
        for prefix, tfix in CULTURE_PREFIX_TYPE:
            if name == prefix or name.startswith(prefix + " "):
                if typ_s != tfix:
                    setc(r, "Type", tfix)
                    typ_s = tfix
                    log.append(f"TYPE culture {name} → {tfix}")
                break
        else:
            if name in EXACT_TYPE and typ_s != EXACT_TYPE[name]:
                setc(r, "Type", EXACT_TYPE[name])
                typ_s = EXACT_TYPE[name]
                log.append(f"TYPE exact {name} → {typ_s}")
            else:
                # substring exact keys (e.g. Adult Red Dragon stays Dragon if typed wrong)
                for key, tfix in EXACT_TYPE.items():
                    if key in name and typ_s and typ_s != tfix and key != name:
                        # only apply for clear undead/fiend words in name
                        if key in ("Shadow", "Vampire", "Ghoul", "Ghost") and key in name.split():
                            setc(r, "Type", tfix)
                            typ_s = tfix
                            log.append(f"TYPE contains {name} → {tfix}")
                            break

        # Dragon name heuristic
        if "Dragon" in name and typ_s not in ("Dragon",):
            setc(r, "Type", "Dragon")
            typ_s = "Dragon"
            log.append(f"TYPE dragon {name}")

        # Giant name heuristic
        if name.endswith("Giant") or " Giant" in name:
            if typ_s != "Giant":
                setc(r, "Type", "Giant")
                typ_s = "Giant"
                log.append(f"TYPE giant {name}")

        # Replace generic attacks
        a1 = get(r, "Attack1 Name")
        if a1 and str(a1) in GENERIC_ATTACKS:
            pack = ATTACKS_BY_TYPE.get(typ_s or "Monstrosity", ATTACKS_BY_TYPE["Monstrosity"])
            setc(r, "Attack1 Name", pack[0])
            setc(r, "Attack1 Type", pack[2])
            a2 = get(r, "Attack2 Name")
            if a2 and str(a2) in GENERIC_ATTACKS:
                setc(r, "Attack2 Name", pack[1])
                setc(r, "Attack2 Type", pack[3])
            # clear trait boilerplate about Medium monstrosity band if present
            traits = get(r, "Traits")
            if traits and "excel-sandbox" in str(traits):
                setc(r, "Traits", f"Provisional combat shell ({typ_s or 'Unknown'}; lore-pass 2026-07-30).")
            setc(r, "Combat Provenance", "lore-pass-2026-07-30-attacks")

        # Size sanity for common animals wrongly Huge
        nr = get(r, "NeedsReview")
        if nr and "cr-suspect-high" in str(nr):
            # light touch: Cow/Swarm should not be Huge CR 10
            if name in ("Cow", "Swarm of Ravens", "Swarm of Bats", "Quipper"):
                setc(r, "Size", "Medium" if "Swarm" not in name else "Medium")
                if name == "Quipper":
                    setc(r, "Size", "Tiny")
                    setc(r, "Challenge", 0)
                    setc(r, "XP", 10)
                    setc(r, "HP", 1)
                    setc(r, "AC", 13)
                if name == "Cow":
                    setc(r, "Size", "Large")
                    setc(r, "Challenge", "1/4")
                    setc(r, "XP", 50)
                    setc(r, "HP", 15)
                    setc(r, "AC", 10)
                if name.startswith("Swarm"):
                    setc(r, "Challenge", "1/4")
                    setc(r, "XP", 50)
                    setc(r, "HP", 24)
                    setc(r, "AC", 12)
                setc(r, "NeedsReview", "corrected-lore-pass-2026-07-30")
                log.append(f"SIZE/CR sanity {name}")

    wb.save(XLSX)
    log_path = ROOT / "docs" / "mechanics" / "BESTIARY_LORE_PASS_2026-07-30.md"
    run_stamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    run_block = (
        f"\n\n## Run log — {run_stamp}\n\n"
        f"**Backup:** `{BACKUP.name}`\n\n"
        f"**Deleted rows:** {deleted}\n\n"
        f"**Log lines:** {len(log)}\n\n"
        + "\n".join(f"- {x}" for x in log[:200])
        + ("\n\n*(truncated)*\n" if len(log) > 200 else "\n")
    )
    if log_path.exists() and "Summary for owner review" in log_path.read_text(encoding="utf-8"):
        # Preserve stamped owner summary; append this run only.
        with log_path.open("a", encoding="utf-8") as f:
            f.write(run_block)
    else:
        log_path.write_text(
            "# Bestiary lore quality pass — 2026-07-30\n\n"
            f"**Backup:** `{BACKUP.name}`\n\n"
            f"**Owner:** sign-off to apply lore-aligned Excel changes.\n\n"
            f"**Deleted rows:** {deleted}\n\n"
            f"**Log lines:** {len(log)}\n\n"
            "## Changes\n\n"
            + "\n".join(f"- {x}" for x in log[:500])
            + ("\n\n*(truncated)*\n" if len(log) > 500 else "\n"),
            encoding="utf-8",
        )
    print("saved", XLSX)
    print("log", log_path, "entries", len(log), "deleted", deleted)


if __name__ == "__main__":
    main()
