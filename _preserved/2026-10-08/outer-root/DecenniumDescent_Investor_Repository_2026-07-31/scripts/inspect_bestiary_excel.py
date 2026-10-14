#!/usr/bin/env python3
"""One-off inspect: Excel columns vs JSON placeholders."""
import json
from pathlib import Path

import openpyxl
import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "bestiary_metadata_source.xlsx"
JSON_PATH = ROOT / "src" / "bestiary_data_full.json"

wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
print("=== SHEETS ===")
for s in wb.sheetnames:
    print(f"  {s}")
wb.close()

xl = pd.ExcelFile(XLSX)
for sheet in xl.sheet_names:
    df = pd.read_excel(XLSX, sheet_name=sheet)
    print(f"\n=== SHEET: {sheet} ===")
    print(f"Rows: {len(df)}")
    print(f"Columns ({len(df.columns)}):")
    for i, c in enumerate(df.columns):
        non_null = int(df[c].notna().sum())
        sample = df[c].dropna().head(2).tolist()
        print(f"  [{i}] {repr(c)}: {non_null} non-null, sample={sample}")

with JSON_PATH.open(encoding="utf-8") as f:
    data = json.load(f)
print(f"\n=== JSON bestiary_data_full.json ===")
print(f"Total entries: {len(data)}")
stub_ac_hp = sum(1 for c in data if c.get("armorClass") == 10 and c.get("hitPoints") == 1)
print(f"AC10/HP1 entries: {stub_ac_hp}")
non_stub = [c for c in data if not (c.get("armorClass") == 10 and c.get("hitPoints") == 1)]
print(f"Non-stub entries: {len(non_stub)}")
if non_stub:
    print("Sample non-stub:")
    for c in non_stub[:5]:
        print(
            f"  {c['name']}: AC={c.get('armorClass')} HP={c.get('hitPoints')} CR={c.get('challengeRating')}"
        )
has_combat = sum(1 for c in data if c.get("combat"))
print(f"Entries with combat payload: {has_combat}")

# Compare Excel names to JSON
monsters_df = pd.read_excel(XLSX, sheet_name=xl.sheet_names[0])
excel_names = set(monsters_df["Name"].dropna().astype(str))
json_names = {c["name"] for c in data}
print(f"\nExcel unique names: {len(excel_names)}")
print(f"JSON unique names: {len(json_names)}")
print(f"In Excel not JSON: {len(excel_names - json_names)}")
print(f"In JSON not Excel: {len(json_names - excel_names)}")
if excel_names - json_names:
    for n in sorted(excel_names - json_names)[:10]:
        print(f"  missing from JSON: {n}")
if json_names - excel_names:
    for n in sorted(json_names - excel_names)[:10]:
        print(f"  extra in JSON: {n}")
