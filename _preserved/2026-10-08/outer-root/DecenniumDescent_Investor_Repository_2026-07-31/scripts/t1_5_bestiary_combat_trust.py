# -*- coding: utf-8 -*-
"""
T1–5 bestiary combat trust pass (2026-07-31).

Raises combat trust for creatures relevant to dungeon Tiers 1–5:
- Zone aliases + CR caps from Director bible seal ladder
- Type-flavored usable Attack1/2 (replaces sandbox / lore-pass generics)
- Freezes T1–5 Gate Guardians (Vaskor, Bloom, Ashthirst, Gor Vekk, Choralith)
  as fully trusted combat + locked metadata
- Metadata QA: Size/Type/Alignment/CR/Zone/Tags; collapse repeated rename prefixes
- Does NOT invent Silent Ones lore; does NOT wipe lore renames
- Idempotent; backs up Bestiary.xlsx first

Usage:
  python scripts/t1_5_bestiary_combat_trust.py
  npm run ingest:bestiary
"""
from __future__ import annotations

import hashlib
import re
import shutil
from datetime import datetime
from pathlib import Path
from typing import Any

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys_path_scripts = ROOT / "scripts"
import sys

sys.path.insert(0, str(sys_path_scripts))
from scrub_bestiary_ip import rename_value  # noqa: E402

XLSX = ROOT / "Bestiary.xlsx"
BACKUP_DIR = ROOT / "backups"
STAMP = datetime.now().strftime("%Y%m%d_%H%M%S")
BACKUP = BACKUP_DIR / f"Bestiary_pre_t15_combat_{STAMP}.xlsx"

PROVENANCE_TRASH = "t1-5-trusted-2026-07-31"
PROVENANCE_GATE = "gate-guardian-trusted-2026-07-31"

# Director bible: floors 1–50 seals. Zone aliases as they appear in Excel.
ZONE_ALIASES: dict[int, list[str]] = {
    1: ["Yieldheart Caverns", "Yieldheart"],
    2: ["Bloomdeep Marsh", "Bloomdeep", "Aethel Swamps"],
    3: ["Anvil of Brogi", "Bastion of Flame"],
    4: ["Isles of Eternal Radiance", "Radiant Isles"],
    5: ["Sanctum of Echoing Choirs", "Choir Sanctum"],
}
# Include trash up to seal CR (locked Gate Guardian ladder).
CR_CAP: dict[int, float] = {1: 8, 2: 11, 3: 14, 4: 17, 5: 20}

# Canonical zone display for clear dungeon aliases only (not cosmology Bastion / surface Swamps).
ZONE_NORMALIZE: dict[str, str] = {
    "yieldheart": "Yieldheart Caverns",
    "yieldheart caverns": "Yieldheart Caverns",
    "bloomdeep": "Bloomdeep Marsh",
    "bloomdeep marsh": "Bloomdeep Marsh",
    "radiant isles": "Isles of Eternal Radiance",
    "isles of eternal radiance": "Isles of Eternal Radiance",
    "choir sanctum": "Sanctum of Echoing Choirs",
    "sanctum of echoing choirs": "Sanctum of Echoing Choirs",
    "anvil of brogi": "Anvil of Brogi",
}

# Locked T1–5 Gate Guardians (CR ladder from lore pass / Director bible).
T15_GUARDIANS: dict[str, dict[str, Any]] = {
    "Vaskor, the Stone Maw": {
        "Size": "Huge",
        "Type": "Elemental",
        "Alignment": "Neutral",
        "Challenge": 8,
        "XP": 3900,
        "Zone Found": "Yieldheart Caverns",
        "Tags": "Gate Guardian",
        "AC": 17,
        "HP": 284,
        "Speed": "30 ft., burrow 20 ft.",
        "Attack1 Name": "Stone Maw",
        "Attack1 Bonus": 7,
        "Attack1 Damage": "2d10+7",
        "Attack1 Type": "bludgeoning",
        "Attack2 Name": "Crushing Jaw",
        "Attack2 Bonus": 7,
        "Attack2 Damage": "2d8+5",
        "Attack2 Type": "bludgeoning",
        "Multiattack": 2,
        "Save DC": 15,
        "Recharge Name": "Seismic Roar",
        "Recharge Kind": "save",
        "Recharge Damage": "5d8",
        "Recharge Type": "bludgeoning",
        "Recharge Save DC": 15,
        "Recharge Save Ability": "strength",
        "Recharge Roll": "5-6",
        "Traits": "Trusted T1 seal — earth-maw elemental; burrow/resurface; Seismic Roar punishes clustered Support.",
    },
    "The Bloom That Hungers": {
        "Size": "Huge",
        "Type": "Plant",
        "Alignment": "Chaotic Evil",
        "Challenge": 11,
        "XP": 7200,
        "Zone Found": "Bloomdeep Marsh",
        "Tags": "Gate Guardian",
        "AC": 15,
        "HP": 374,
        "Speed": "20 ft., burrow 30 ft.",
        "Attack1 Name": "Vine Maw",
        "Attack1 Bonus": 8,
        "Attack1 Damage": "2d10+9",
        "Attack1 Type": "piercing",
        "Attack2 Name": "Root Constrict",
        "Attack2 Bonus": 8,
        "Attack2 Damage": "2d8+6",
        "Attack2 Type": "bludgeoning",
        "Multiattack": 2,
        "Save DC": 16,
        "Recharge Name": "Spore Burst",
        "Recharge Kind": "save",
        "Recharge Damage": "4d8",
        "Recharge Type": "poison",
        "Recharge Save DC": 16,
        "Recharge Save Ability": "constitution",
        "Recharge Roll": "5-6",
        "Traits": "Trusted T2 seal — spore-hunger plant; area poison; Vine Maw + Root Constrict attrition.",
    },
    "Ashthirst the Hollow": {
        "Size": "Huge",
        "Type": "Elemental",
        "Alignment": "Neutral Evil",
        "Challenge": 14,
        "XP": 11500,
        "Zone Found": "Anvil of Brogi",
        "Tags": "Gate Guardian",
        "AC": 16,
        "HP": 464,
        "Speed": "40 ft., fly 60 ft. (hover)",
        "Attack1 Name": "Ash Claw",
        "Attack1 Bonus": 9,
        "Attack1 Damage": "3d8+7",
        "Attack1 Type": "necrotic",
        "Attack2 Name": "Cinder Lash",
        "Attack2 Bonus": 9,
        "Attack2 Damage": "2d10+5",
        "Attack2 Type": "fire",
        "Multiattack": 2,
        "Save DC": 18,
        "Recharge Name": "Hollow Wail",
        "Recharge Kind": "save",
        "Recharge Damage": "6d8",
        "Recharge Type": "psychic",
        "Recharge Save DC": 18,
        "Recharge Save Ability": "wisdom",
        "Recharge Roll": "5-6",
        "Traits": "Trusted T3 seal — hollow forge-elemental; hover over lava; Hollow Wail vs clustered parties.",
    },
    "Gor Vekk, Architect of Ends": {
        "Size": "Large",
        "Type": "Aberration",
        "Alignment": "Lawful Evil",
        "Challenge": 17,
        "XP": 18000,
        "Zone Found": "Isles of Eternal Radiance",
        "Tags": "Gate Guardian",
        "AC": 19,
        "HP": 416,
        "Speed": "50 ft., fly 120 ft.",
        "Attack1 Name": "Reality Sunder",
        "Attack1 Bonus": 10,
        "Attack1 Damage": "3d10+7",
        "Attack1 Type": "force",
        "Attack2 Name": "End Thesis",
        "Attack2 Bonus": 10,
        "Attack2 Damage": "2d10+6",
        "Attack2 Type": "force",
        "Multiattack": 3,
        "Save DC": 19,
        "Recharge Name": "End Architecture",
        "Recharge Kind": "save",
        "Recharge Damage": "7d8",
        "Recharge Type": "force",
        "Recharge Save DC": 19,
        "Recharge Save Ability": "intelligence",
        "Recharge Roll": "5-6",
        "Traits": "Trusted T4 seal — CR 17 seal budget (not CR 29 leftover). Projection fights may use reduced HP.",
    },
    "Choralith, Voice Below": {
        "Size": "Large",
        "Type": "Aberration",
        "Alignment": "Neutral",
        "Challenge": 20,
        "XP": 25000,
        "Zone Found": "Sanctum of Echoing Choirs",
        "Tags": "Gate Guardian",
        "AC": 19,
        "HP": 483,
        "Speed": "30 ft., fly 60 ft.",
        "Attack1 Name": "Resonant Lash",
        "Attack1 Bonus": 11,
        "Attack1 Damage": "3d12+8",
        "Attack1 Type": "thunder",
        "Attack2 Name": "Echo Needle",
        "Attack2 Bonus": 11,
        "Attack2 Damage": "2d10+6",
        "Attack2 Type": "psychic",
        "Multiattack": 2,
        "Save DC": 20,
        "Recharge Name": "Subsonic Chord",
        "Recharge Kind": "save",
        "Recharge Damage": "6d10",
        "Recharge Type": "thunder",
        "Recharge Save DC": 20,
        "Recharge Save Ability": "constitution",
        "Recharge Roll": "5-6",
        "Traits": "Trusted T5 seal — echo/morale pressure; soft quit at Canticle is morale, not Opening Horn.",
    },
}

# DMG-style CR bands (same spine as generate_bestiary_combat_sandbox.py).
CR_BANDS: dict[str, dict[str, Any]] = {
    "0": {"ac": 10, "hp": 5, "bonus": 3, "dice": "1d4", "mod": 0, "multi": 1, "save_dc": None},
    "1/8": {"ac": 11, "hp": 7, "bonus": 3, "dice": "1d4", "mod": 1, "multi": 1, "save_dc": None},
    "1/4": {"ac": 11, "hp": 11, "bonus": 3, "dice": "1d6", "mod": 1, "multi": 1, "save_dc": None},
    "1/2": {"ac": 12, "hp": 22, "bonus": 4, "dice": "1d6", "mod": 2, "multi": 1, "save_dc": None},
    "1": {"ac": 13, "hp": 35, "bonus": 4, "dice": "1d8", "mod": 2, "multi": 1, "save_dc": None},
    "2": {"ac": 13, "hp": 45, "bonus": 5, "dice": "1d10", "mod": 2, "multi": 1, "save_dc": None},
    "3": {"ac": 14, "hp": 58, "bonus": 5, "dice": "2d6", "mod": 2, "multi": 1, "save_dc": 11},
    "4": {"ac": 14, "hp": 75, "bonus": 6, "dice": "2d6", "mod": 3, "multi": 1, "save_dc": 12},
    "5": {"ac": 15, "hp": 95, "bonus": 6, "dice": "2d8", "mod": 4, "multi": 1, "save_dc": 13},
    "6": {"ac": 15, "hp": 112, "bonus": 6, "dice": "2d8", "mod": 4, "multi": 2, "save_dc": 13},
    "7": {"ac": 15, "hp": 123, "bonus": 7, "dice": "2d8", "mod": 5, "multi": 2, "save_dc": 14},
    "8": {"ac": 16, "hp": 142, "bonus": 7, "dice": "2d10", "mod": 5, "multi": 2, "save_dc": 15},
    "9": {"ac": 16, "hp": 157, "bonus": 7, "dice": "2d10", "mod": 5, "multi": 2, "save_dc": 15},
    "10": {"ac": 17, "hp": 172, "bonus": 8, "dice": "2d10", "mod": 6, "multi": 2, "save_dc": 16},
    "11": {"ac": 17, "hp": 187, "bonus": 8, "dice": "2d10", "mod": 6, "multi": 2, "save_dc": 16},
    "12": {"ac": 17, "hp": 202, "bonus": 8, "dice": "3d8", "mod": 6, "multi": 2, "save_dc": 17},
    "13": {"ac": 18, "hp": 217, "bonus": 8, "dice": "3d8", "mod": 6, "multi": 2, "save_dc": 17},
    "14": {"ac": 18, "hp": 232, "bonus": 9, "dice": "3d8", "mod": 7, "multi": 2, "save_dc": 18},
    "15": {"ac": 18, "hp": 247, "bonus": 9, "dice": "3d10", "mod": 7, "multi": 2, "save_dc": 18},
    "16": {"ac": 18, "hp": 262, "bonus": 9, "dice": "3d10", "mod": 7, "multi": 3, "save_dc": 19},
    "17": {"ac": 19, "hp": 277, "bonus": 10, "dice": "3d10", "mod": 7, "multi": 3, "save_dc": 19},
    "18": {"ac": 19, "hp": 292, "bonus": 10, "dice": "3d10", "mod": 7, "multi": 3, "save_dc": 19},
    "19": {"ac": 19, "hp": 307, "bonus": 10, "dice": "3d12", "mod": 8, "multi": 3, "save_dc": 20},
    "20": {"ac": 19, "hp": 322, "bonus": 11, "dice": "3d12", "mod": 8, "multi": 3, "save_dc": 20},
}

SIZE_MODS: dict[str, dict[str, float | int]] = {
    "Tiny": {"ac": 1, "hp_mult": 0.25, "dmg_mult": 0.5, "multi_add": 0},
    "Small": {"ac": 0, "hp_mult": 0.5, "dmg_mult": 0.75, "multi_add": 0},
    "Medium": {"ac": 0, "hp_mult": 1.0, "dmg_mult": 1.0, "multi_add": 0},
    "Large": {"ac": -1, "hp_mult": 1.5, "dmg_mult": 1.25, "multi_add": 0},
    "Huge": {"ac": -2, "hp_mult": 2.0, "dmg_mult": 1.5, "multi_add": 1},
    "Gargantuan": {"ac": -2, "hp_mult": 2.5, "dmg_mult": 2.0, "multi_add": 1},
}

TYPE_ATTACKS: dict[str, list[tuple[str, str]]] = {
    "aberration": [("Mind Lash", "psychic"), ("Warped Tentacle", "bludgeoning"), ("Void Rend", "psychic")],
    "beast": [("Rending Claw", "slashing"), ("Crushing Bite", "piercing"), ("Goring Horn", "piercing")],
    "celestial": [("Radiant Smite", "radiant"), ("Luminous Strike", "radiant"), ("Sanctified Blow", "radiant")],
    "construct": [("Iron Slam", "bludgeoning"), ("Stone Fist", "bludgeoning"), ("Hammer Blow", "bludgeoning")],
    "dragon": [("Scale Rend", "slashing"), ("Drake Claw", "slashing"), ("Wyrm Bite", "piercing")],
    "elemental": [("Primal Strike", "bludgeoning"), ("Elemental Slam", "thunder"), ("Core Burst", "fire")],
    "fey": [("Glamour Cut", "psychic"), ("Wild Thorn", "piercing"), ("Fey Whisper", "psychic")],
    "fiend": [("Hellfire Claw", "fire"), ("Abyssal Rend", "necrotic"), ("Infernal Strike", "fire")],
    "giant": [("Massive Club", "bludgeoning"), ("Boulder Smash", "bludgeoning"), ("Titan Fist", "bludgeoning")],
    "humanoid": [("Weapon Strike", "slashing"), ("Heavy Blow", "bludgeoning"), ("Precise Cut", "slashing")],
    "monstrosity": [("Hooked Claw", "slashing"), ("Crushing Maw", "piercing"), ("Barbed Lash", "slashing")],
    "ooze": [("Acid Pseudopod", "acid"), ("Corrosive Touch", "acid"), ("Engulfing Slam", "acid")],
    "plant": [("Thorn Lash", "piercing"), ("Root Strike", "bludgeoning"), ("Spore Touch", "poison")],
    "undead": [("Grave Touch", "necrotic"), ("Bone Rake", "slashing"), ("Death Chill", "cold")],
}

TYPE_ATTACKS_SECONDARY: dict[str, list[tuple[str, str]]] = {
    "aberration": [("Psychic Pulse", "psychic")],
    "beast": [("Tail Swipe", "bludgeoning")],
    "celestial": [("Holy Bolt", "radiant")],
    "construct": [("Grinding Gear", "bludgeoning")],
    "dragon": [("Tail Sweep", "bludgeoning")],
    "elemental": [("Surge Bolt", "lightning")],
    "fey": [("Mislead Strike", "psychic")],
    "fiend": [("Soul Burn", "fire")],
    "giant": [("Ground Stomp", "bludgeoning")],
    "humanoid": [("Off-hand Strike", "piercing")],
    "monstrosity": [("Tail Lash", "bludgeoning")],
    "ooze": [("Dissolve", "acid")],
    "plant": [("Entangle", "bludgeoning")],
    "undead": [("Life Drain", "necrotic")],
}

SPEED_BY_TYPE: dict[str, str] = {
    "aberration": "30 ft.",
    "beast": "40 ft.",
    "celestial": "30 ft., fly 60 ft.",
    "construct": "30 ft.",
    "dragon": "40 ft., fly 80 ft.",
    "elemental": "40 ft.",
    "fey": "30 ft., fly 30 ft.",
    "fiend": "30 ft., fly 60 ft.",
    "giant": "40 ft.",
    "humanoid": "30 ft.",
    "monstrosity": "40 ft.",
    "ooze": "20 ft., climb 20 ft.",
    "plant": "20 ft.",
    "undead": "30 ft.",
}

# Lore-pass / sandbox generics that should be replaced for the T1–5 set.
REPLACEABLE_ATTACKS = {
    "Freak Bite",
    "Monstrous Claw",
    "Aberrant Lash",
    "Tail Lash",
    "Grave Shriek",
    "War Blade",
    "Crossbow Bolt",
    "Hellish Claw",
    "Infernal Lash",
    "Tree-Trunk Club",
    "Boulder Hurl",
    "Iron Fist",
    "Gear Shear",
    "Elemental Slam",
    "Surge Burst",
    "Vine Lash",
    "Hooked Claw",
    "Crushing Maw",
    "Warp Tentacle",
    "Mind Spike",
    "Glamour Blade",
    "Hex Dart",
    "Grave Touch",
    "Soul Chill",
    "Drake Bite",
    "Tail Sweep",
    "Rending Bite",
    "Claw Swipe",
    "Radiant Smite",
    "Judgment Beam",
    "Acid Pseudopod",
    "Engulf Pulse",
    "Spore Burst",  # only as Attack1/2 at-will placeholder; guardians keep recharge name
}

REPEATED_PREFIXES = ("Fracture-Kin", "Order-Drone", "Hush-Broker")

SIZE_ORDER = ["Tiny", "Small", "Medium", "Large", "Huge", "Gargantuan"]

# Culture-prefix and archetype Type floors (from lore quality pass; no Silent Ones inventing).
CULTURE_PREFIX_TYPE = [
    ("Umbrin", "Humanoid"),
    ("Deep Dwarf", "Humanoid"),
    ("Serpent-Blood", "Monstrosity"),
    ("Astral Reaver", "Humanoid"),
    ("Astral Ascetic", "Humanoid"),
    ("Depth Toad", "Humanoid"),
    ("Gaze Tyrant", "Aberration"),
    ("Voidmind", "Aberration"),
    ("Sporefolk", "Plant"),
    ("Cave Gnome", "Humanoid"),
    ("Poison-Skin", "Humanoid"),
    ("Shellkin", "Humanoid"),
    ("Order-Drone", "Construct"),
    ("Hush-Broker", "Fiend"),
    ("Ferry-Wraith", "Fiend"),
    ("Court Broker", "Fiend"),
    ("Fracture-Kin", "Fiend"),
    ("Pact Devil", "Fiend"),
    ("Chaos-Husk", "Aberration"),
    ("Season-Bound", "Fey"),
    ("Grief-Bound", "Monstrosity"),
    ("Spiral Breach", "Aberration"),
    ("Void Broker", "Aberration"),
]

EXACT_TYPE = {
    "Awakened Tree": "Plant",
    "Assassin": "Humanoid",
    "Allip": "Undead",
    "Bodak": "Undead",
    "Ghost": "Undead",
    "Ghoul": "Undead",
    "Specter": "Undead",
    "Wraith": "Undead",
    "Revenant": "Undead",
    "Vampire": "Undead",
    "Mummy": "Undead",
    "Zombie": "Undead",
    "Skeleton": "Undead",
    "Deathlock": "Undead",
    "Flameskull": "Undead",
    "Will-o-wisp": "Undead",
    "Dybbuk": "Fiend",
    "Aboleth": "Aberration",
    "Otyugh": "Aberration",
    "Cloaker": "Aberration",
    "Gibbering Mouther": "Aberration",
    "Berbalang": "Aberration",
    "Dimetrodon": "Beast",
    "Panther": "Beast",
    "Giant Boar": "Beast",
    "Giant Rat": "Beast",
    "Giant Toad": "Beast",
    "Giant Centipede": "Beast",
    "Giant Hyena": "Beast",
    "Riding Horse": "Beast",
    "Sea Horse": "Beast",
    "Octopus": "Beast",
    "Scorpion": "Beast",
    "Stegosaurus": "Beast",
    "Cockatrice": "Monstrosity",
    "Peryton": "Monstrosity",
    "Manticore": "Monstrosity",
    "Behir": "Monstrosity",
    "Griffon": "Monstrosity",
    "Harpy": "Monstrosity",
    "Banderhobb": "Monstrosity",
    "Gray Ooze": "Ooze",
    "Mud Mephit": "Elemental",
    "Dust Mephit": "Elemental",
    "Smoke Mephit": "Elemental",
    "Invisible Stalker": "Elemental",
    "Fire Elemental Myrmidon": "Elemental",
    "Treant": "Plant",
    "Shambling Mound": "Plant",
    "Redcap": "Fey",
    "Korred": "Fey",
    "Darkling Elder": "Fey",
    "Nilbog": "Humanoid",
    "Knight": "Humanoid",
    "Scout": "Humanoid",
    "Warlord": "Humanoid",
    "Master Thief": "Humanoid",
    "Druid": "Humanoid",
    "Illusionist": "Humanoid",
    "Conjurer": "Humanoid",
    "Transmuter": "Humanoid",
    "Orc": "Humanoid",
    "Orog": "Humanoid",
    "Centaur": "Monstrosity",
    "Half-ogre": "Giant",
    "Stone Giant": "Giant",
    "Stone Giant Dreamwalker": "Giant",
    "Fire Giant Dreadnought": "Giant",
    "Dire Troll": "Giant",
    "Ogre Battering Ram": "Giant",
    "Cambion": "Fiend",
    "Lemure": "Fiend",
    "Spined Devil": "Fiend",
    "Deva": "Celestial",
    "Lizardfolk": "Humanoid",
    "Kobold Scale Sorcerer": "Humanoid",
    "Firenewt Warrior": "Humanoid",
    "Sahuagin Priestess": "Humanoid",
    "Hobgoblin Devastator": "Humanoid",
    "Deep Scion": "Humanoid",
    "Shadow Dancer": "Humanoid",
    "Retriever": "Construct",
    "Steel Predator": "Construct",
    "Oni": "Giant",
    "Devourer": "Undead",
    "Kraken": "Monstrosity",
    "Swarm of Cranium Rats": "Beast",
    "Umbrin Arachnid": "Monstrosity",
    "Umbrin Shadowblade": "Humanoid",
    "Spiral Worm": "Aberration",
    "Corrosion Grazer": "Monstrosity",
    "Shellkin": "Humanoid",
}


def infer_type(name: str, current: str | None) -> str | None:
    if name in EXACT_TYPE:
        return EXACT_TYPE[name]
    for prefix, typ in CULTURE_PREFIX_TYPE:
        if name.startswith(prefix):
            return typ
    # Soft tokens
    low = name.lower()
    if "dragon" in low and "half-" not in low:
        return "Dragon"
    if low.endswith(" golem") or "golem" in low:
        return "Construct"
    if "mephit" in low:
        return "Elemental"
    if current and str(current).strip():
        return None  # keep existing when no rule
    return "Monstrosity"


def normalize_cr(raw: Any) -> str:
    if raw is None:
        return "0"
    text = str(raw).strip()
    if not text or text.lower() == "nan":
        return "0"
    if text.endswith(".0"):
        text = text[:-2]
    return text


def cr_to_float(cr: str) -> float:
    if "/" in cr:
        num, den = cr.split("/", 1)
        return float(num) / float(den)
    return float(cr)


def match_tier(zone: str | None) -> int | None:
    z = (zone or "").strip().lower()
    if not z:
        return None
    for tier, aliases in ZONE_ALIASES.items():
        for alias in aliases:
            a = alias.lower()
            if z == a or z.startswith(a) or a.startswith(z):
                return tier
    return None


def pick_from_pool(name: str, pool: list[tuple[str, str]]) -> tuple[str, str]:
    idx = int(hashlib.md5(name.encode("utf-8")).hexdigest(), 16) % len(pool)
    return pool[idx]


def format_damage(dice: str, mod: int) -> str:
    if mod > 0:
        return f"{dice}+{mod}"
    if mod < 0:
        return f"{dice}{mod}"
    return dice


def type_key(ctype: Any) -> str:
    key = str(ctype or "monstrosity").lower().split("(")[0].strip()
    return key if key in TYPE_ATTACKS else "monstrosity"


def size_floor_for_cr(cr_f: float) -> str:
    if cr_f <= 0.25:
        return "Tiny"
    if cr_f < 1:
        return "Small"
    if cr_f < 5:
        return "Medium"
    if cr_f < 11:
        return "Large"
    if cr_f < 20:
        return "Huge"
    return "Gargantuan"


def collapse_repeated_prefix(name: str) -> str:
    out = name.strip()
    for prefix in REPEATED_PREFIXES:
        token = prefix + " "
        while out.startswith(token + prefix):
            out = out[len(token) :]
    return out


def colmap(ws) -> dict[str, int]:
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    return {str(h): i + 1 for i, h in enumerate(header) if h}


def ensure_cols(ws, cols: dict[str, int], names: list[str]) -> dict[str, int]:
    for name in names:
        if name in cols:
            continue
        new_i = max(cols.values(), default=0) + 1
        ws.cell(1, new_i, name)
        cols[name] = new_i
    return cols


def generate_trash_combat(name: str, size: str, ctype: str, cr: str) -> dict[str, Any]:
    band = CR_BANDS.get(cr, CR_BANDS["0"])
    size_key = size if size in SIZE_MODS else "Medium"
    sm = SIZE_MODS[size_key]
    tkey = type_key(ctype)
    cr_f = cr_to_float(cr)

    ac = max(8, int(band["ac"]) + int(sm["ac"]))
    hp = max(1, round(int(band["hp"]) * float(sm["hp_mult"])))
    bonus = int(band["bonus"])
    mod = max(0, round(int(band["mod"]) * float(sm["dmg_mult"])))
    multi = max(1, int(band["multi"]) + int(sm["multi_add"]))
    if cr_f >= 11:
        multi = max(multi, 2)

    atk1_name, atk1_type = pick_from_pool(name, TYPE_ATTACKS[tkey])
    atk1_damage = format_damage(str(band["dice"]), mod)

    out: dict[str, Any] = {
        "AC": ac,
        "HP": hp,
        "Speed": SPEED_BY_TYPE.get(tkey, "30 ft."),
        "Attack1 Name": atk1_name,
        "Attack1 Bonus": bonus,
        "Attack1 Damage": atk1_damage,
        "Attack1 Type": atk1_type,
        "Attack2 Name": None,
        "Attack2 Bonus": None,
        "Attack2 Damage": None,
        "Attack2 Type": None,
        "Multiattack": multi if multi > 1 else None,
        "Save DC": band.get("save_dc"),
        "Traits": f"Trusted T1–5 usable draft — CR {cr} {size_key} {tkey} band (type-flavored).",
        "Combat Provenance": PROVENANCE_TRASH,
    }
    if multi >= 2 or cr_f >= 5:
        atk2_name, atk2_type = pick_from_pool(name + "_2", TYPE_ATTACKS_SECONDARY.get(tkey, TYPE_ATTACKS[tkey]))
        out["Attack2 Name"] = atk2_name
        out["Attack2 Bonus"] = bonus
        out["Attack2 Damage"] = format_damage(str(band["dice"]), max(0, mod - 1))
        out["Attack2 Type"] = atk2_type
    return out


def already_trusted(prov: Any) -> bool:
    p = str(prov or "").lower()
    return "t1-5-trusted-2026-07-31" in p or "gate-guardian-trusted-2026-07-31" in p


def main() -> None:
    if not XLSX.exists():
        raise SystemExit(f"Missing {XLSX}")

    BACKUP_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(XLSX, BACKUP)
    print("backup", BACKUP)

    wb = openpyxl.load_workbook(XLSX)
    ws = wb["Monsters"]
    cols = colmap(ws)
    cols = ensure_cols(
        ws,
        cols,
        [
            "Description",
            "NeedsReview",
            "Combat Provenance",
            "Traits",
            "Recharge Name",
            "Recharge Kind",
            "Recharge Bonus",
            "Recharge Damage",
            "Recharge Type",
            "Recharge Save DC",
            "Recharge Save Ability",
            "Recharge Roll",
        ],
    )

    def get(row: int, key: str, default=None):
        i = cols.get(key)
        if not i:
            return default
        return ws.cell(row, i).value

    def setc(row: int, key: str, value):
        i = cols.get(key)
        if not i:
            return
        ws.cell(row, i).value = value

    upgraded_trash = 0
    frozen_guardians = 0
    meta_fixes = 0
    skipped_over_cap = 0
    selected_names: list[tuple[int, str]] = []
    sandbox_outside = 0

    # Pass 1: freeze T1–5 guardians (always, idempotent overwrite of combat lock fields).
    for r in range(2, ws.max_row + 1):
        name = str(get(r, "Name") or "").strip()
        if name not in T15_GUARDIANS:
            continue
        spec = T15_GUARDIANS[name]
        for k, v in spec.items():
            setc(r, k, v)
        setc(r, "Combat Provenance", PROVENANCE_GATE)
        setc(r, "NeedsReview", None)
        frozen_guardians += 1
        tier = match_tier(str(spec["Zone Found"])) or 0
        selected_names.append((tier, name))
        print(f"GATE freeze {name}: CR {spec['Challenge']} AC {spec['AC']} HP {spec['HP']}")

    # Pass 2: T1–5 trash by zone + CR cap.
    for r in range(2, ws.max_row + 1):
        name = str(get(r, "Name") or "").strip()
        if not name or name in T15_GUARDIANS:
            continue

        # In-memory canonical name for Type rules (Excel stays pre-scrub; ingest renames on read).
        canon_name = collapse_repeated_prefix(rename_value(name))

        zone = str(get(r, "Zone Found") or "").strip()
        tier = match_tier(zone)
        if tier is None:
            prov = str(get(r, "Combat Provenance") or "")
            if "excel-sandbox" in prov.lower() or "lore-pass-2026-07-30" in prov.lower() or not prov:
                sandbox_outside += 1
            continue

        cr = normalize_cr(get(r, "Challenge"))
        try:
            cr_f = cr_to_float(cr)
        except ValueError:
            cr_f = 0.0

        if cr_f > CR_CAP[tier] + 0.001:
            # Leave outside T1–5 trust; flag for Director review.
            notes = str(get(r, "NeedsReview") or "")
            flag = f"t15-over-cap:T{tier}-CR{cr}"
            if flag not in notes:
                setc(r, "NeedsReview", (notes + "; " + flag).strip("; ").strip())
                meta_fixes += 1
            skipped_over_cap += 1
            continue

        selected_names.append((tier, canon_name))

        # Metadata QA — collapse scrub double-prefix in Excel name only when present.
        collapsed = collapse_repeated_prefix(name)
        if collapsed != name:
            setc(r, "Name", collapsed)
            name = collapsed
            canon_name = collapse_repeated_prefix(rename_value(name))
            meta_fixes += 1

        size = str(get(r, "Size") or "Medium").strip() or "Medium"
        ctype = str(get(r, "Type") or "Monstrosity").strip() or "Monstrosity"
        if size not in SIZE_ORDER:
            size = "Medium"
            setc(r, "Size", size)
            meta_fixes += 1

        floor = size_floor_for_cr(cr_f)
        si = SIZE_ORDER.index(size)
        fi = SIZE_ORDER.index(floor)
        if abs(si - fi) >= 2:
            # Snap toward CR-band floor (Worker 2 rule); never shrink guardians (already skipped).
            setc(r, "Size", floor)
            size = floor
            notes = str(get(r, "NeedsReview") or "")
            flag = f"size-snap:{SIZE_ORDER[si]}->{floor}"
            if flag not in notes:
                setc(r, "NeedsReview", (notes + "; " + flag).strip("; ").strip() if notes else flag)
            meta_fixes += 1

        align = get(r, "Alignment")
        if align is None or str(align).strip() == "":
            setc(r, "Alignment", "Unaligned")
            meta_fixes += 1

        zkey = zone.lower()
        if zkey in ZONE_NORMALIZE and ZONE_NORMALIZE[zkey] != zone:
            setc(r, "Zone Found", ZONE_NORMALIZE[zkey])
            meta_fixes += 1

        # Type metadata QA before combat flavoring (use canonical lore name).
        inferred = infer_type(canon_name, ctype)
        type_changed = False
        if inferred and inferred != ctype:
            setc(r, "Type", inferred)
            ctype = inferred
            type_changed = True
            meta_fixes += 1

        # Combat upgrade: skip when already trusted unless Type just changed.
        prov = get(r, "Combat Provenance")
        a1 = str(get(r, "Attack1 Name") or "").strip()
        if already_trusted(prov) and not type_changed:
            continue
        needs_combat = (not a1) or (a1 in REPLACEABLE_ATTACKS) or (not already_trusted(prov))

        if needs_combat:
            combat = generate_trash_combat(canon_name, size, ctype, cr)
            # Preserve HP/AC floors if existing numbers are already at/above band (Worker 2 bumps).
            try:
                old_hp = int(get(r, "HP") or 0)
            except (TypeError, ValueError):
                old_hp = 0
            try:
                old_ac = int(get(r, "AC") or 0)
            except (TypeError, ValueError):
                old_ac = 0
            combat["HP"] = max(old_hp, int(combat["HP"]))
            combat["AC"] = max(old_ac, int(combat["AC"]))
            # Keep non-generic existing speed if present.
            old_speed = str(get(r, "Speed") or "").strip()
            if old_speed and "ft" in old_speed.lower():
                combat["Speed"] = old_speed

            for k, v in combat.items():
                setc(r, k, v)
            upgraded_trash += 1
        else:
            setc(r, "Combat Provenance", PROVENANCE_TRASH)
            upgraded_trash += 1

    wb.save(XLSX)
    by_tier: dict[int, int] = {}
    for t, _ in selected_names:
        by_tier[t] = by_tier.get(t, 0) + 1
    summary = (
        f"done guardians={frozen_guardians} trash_upgraded={upgraded_trash} "
        f"meta_fixes={meta_fixes} over_cap_skipped={skipped_over_cap} "
        f"selected={len(selected_names)} sandbox_outside_t15={sandbox_outside} "
        f"by_tier={by_tier}"
    )
    print(summary.encode("ascii", "replace").decode("ascii"))


if __name__ == "__main__":
    main()
