#!/usr/bin/env python3
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'books/player/Decennium_Descent_Quickstart_Reference_v1.docx'
DARK='1D2733'; GOLD='B8860B'; RED='8B2E2E'; PALE='F7F1E3'; GRAY='59636E'

def shade(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=tcPr.find(qn('w:shd'))
    if shd is None: shd=OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'),fill)

def margins(cell, v=80):
    tcPr=cell._tc.get_or_add_tcPr(); tcMar=tcPr.first_child_found_in('w:tcMar')
    if tcMar is None: tcMar=OxmlElement('w:tcMar'); tcPr.append(tcMar)
    for name in ('top','start','bottom','end'):
        node=OxmlElement(f'w:{name}'); node.set(qn('w:w'),str(v)); node.set(qn('w:type'),'dxa'); tcMar.append(node)

def add_title(cell, text):
    p=cell.add_paragraph(); p.paragraph_format.space_before=Pt(2); p.paragraph_format.space_after=Pt(2)
    r=p.add_run(text); r.bold=True; r.font.name='Aptos Display'; r.font.size=Pt(10.5); r.font.color.rgb=RGBColor.from_string(RED)

def bullet(cell, text):
    p=cell.add_paragraph(style='List Bullet'); p.paragraph_format.left_indent=Inches(.16); p.paragraph_format.first_line_indent=Inches(-.10); p.paragraph_format.space_after=Pt(1)
    r=p.add_run(text); r.font.name='Aptos'; r.font.size=Pt(7.4)

def para(cell, text, bold_prefix=None):
    p=cell.add_paragraph(); p.paragraph_format.space_after=Pt(2)
    if bold_prefix and text.startswith(bold_prefix):
        r=p.add_run(bold_prefix); r.bold=True; r.font.size=Pt(7.5)
        r=p.add_run(text[len(bold_prefix):]); r.font.size=Pt(7.5)
    else:
        r=p.add_run(text); r.font.size=Pt(7.5)
    return p

doc=Document(); sec=doc.sections[0]
sec.orientation=WD_ORIENT.LANDSCAPE; sec.page_width=Inches(11); sec.page_height=Inches(8.5)
sec.top_margin=Inches(.35); sec.bottom_margin=Inches(.35); sec.left_margin=Inches(.4); sec.right_margin=Inches(.4)
styles=doc.styles; styles['Normal'].font.name='Aptos'; styles['Normal'].font.size=Pt(7.5)
styles['Normal'].paragraph_format.space_after=Pt(1)

p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(1)
r=p.add_run('DECENNIUM DESCENT'); r.bold=True; r.font.name='Aptos Display'; r.font.size=Pt(18); r.font.color.rgb=RGBColor.from_string(DARK)
r=p.add_run('  |  LEVEL 1 QUICKSTART'); r.bold=True; r.font.name='Aptos Display'; r.font.size=Pt(15); r.font.color.rgb=RGBColor.from_string(RED)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(4)
r=p.add_run('Objective → Finance → Roster → Load → Watches → Encounter → Cargo → Extraction → Haven → Advancement'); r.bold=True; r.font.size=Pt(8.5); r.font.color.rgb=RGBColor.from_string(GOLD)

table=doc.add_table(rows=1, cols=4); table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.autofit=False
widths=[2.52,2.52,2.52,2.52]
for i,c in enumerate(table.rows[0].cells):
    c.width=Inches(widths[i]); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP; margins(c); shade(c, PALE)

c=table.cell(0,0); c.text=''; add_title(c,'YOUR TURN')
for x in ['Move up to Speed.','Take one Action.','Bonus Action only when granted.','One free object interaction.','One Reaction before your next turn.','Only one additional full Action per round from all sources.']: bullet(c,x)
add_title(c,'THE ROLL'); para(c,'d20 + ability modifier + proficiency when trained + circumstances')
for x in ['Critical Success: DC +10.','Success: meet DC.','Failure: below DC.','Critical Failure: DC -10.','Natural 20 improves one degree; natural 1 worsens one.']: bullet(c,x)
add_title(c,'RETREAT'); para(c,'Three Escape successes before three Pursuit successes. Normally only one Escape source per round; a Critical Success may grant two.')

c=table.cell(0,1); c.text=''; add_title(c,'ONE 8-HOUR WATCH')
for x in ['Assign every member to a team or Rest.','Spend one-third of daily Supply.','Resolve team activities.','Add route Progress.','Apply Pressure changes.','Roll 1d6; at or below Pressure causes a complication.','Update cargo, Strain, Wounds, Integrity, and Rest Debt.','Advance the clock eight hours.']: bullet(c,x)
add_title(c,'REST DEBT'); para(c,'Lost or interrupted Rest: mark Due in 2 Watches. Complete one full Rest before the second following Watch ends. Missing it causes 1 Strain and debt remains.')
add_title(c,'PRESSURE 0–6'); para(c,'Pressure 6 causes an automatic Severe complication. Attention cargo adds Pressure only when first exposed, re-exposed, or concealment is lost.')

c=table.cell(0,2); c.text=''; add_title(c,'CARGO & EXTRACTION')
for x in ['Record Load, Extraction Class 0–6, priority 1–5, value, and tags.','Risk or abandon lowest-priority lots first.','Extraction declaration: Pressure + floor(highest EC ÷ 2), max +3.','Test DC: tier DC + highest EC + handling, pursuit, and route modifiers.','Cargo-and-handling modifier caps at +7.','On Failure, Director identifies risk; players choose time/complication or the stated cargo risk.']: bullet(c,x)
add_title(c,'UNKNOWN ROUTE'); para(c,'Map Route is safest. Forced March: Severe route check; Success gains 1 Progress, Pressure +2, and one extra Supply per active member. Failure causes a complication and no Progress.')
add_title(c,'TIER I DEFAULTS');
for x in ['Six-person company.','Eight days of Supply.','Reserve 40% Load for cargo/casualties.','Abort: two Critical members, <2 days Supply, transport Integrity 0, or return deadline.']: bullet(c,x)

c=table.cell(0,3); c.text=''; add_title(c,'TIER I HAVEN FAST PATH')
para(c,'Inspection → Appraisal → Sale → Sponsor Settlement → Recovery')
para(c,'Use the full ten-step procedure only when a disputed claim, strategic cargo, lien, quarantine crisis, contraband case, or faction hearing requires it.')
add_title(c,'THREE MONEY BOXES')
para(c,'Gross cargo:', 'Gross cargo:'); para(c,'Value before taxes and shares.')
para(c,'Net proceeds:', 'Net proceeds:'); para(c,'After mandatory Haven costs and sponsor share.')
para(c,'True profit:', 'True profit:'); para(c,'After wages, repairs, replacements, debts, and reserve.')
add_title(c,'AT SESSION END')
for x in ['Opening hour and Watch.','Location, route Progress, and Pressure.','Supply, camp, and transport Integrity.','Members: team, Rest Debt, Strain, Wounds, HP, RP/RC.','Cargo manifest and priorities.','Objective, abort conditions, and faction changes.']: bullet(c,x)

# Footer
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(3); p.paragraph_format.space_after=Pt(0)
r=p.add_run('V7.2 public reference • Full rules: Gameplay Loop Bible v1.2'); r.font.size=Pt(6.8); r.font.color.rgb=RGBColor.from_string(GRAY)
OUT.parent.mkdir(parents=True,exist_ok=True); doc.save(OUT); print(OUT)
