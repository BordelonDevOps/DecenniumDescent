#!/usr/bin/env python3
from __future__ import annotations
from pathlib import Path
import re
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "books/player/Decennium_Descent_Gameplay_Loop_Bible_v1_1.docx"
SOURCES = [
    ROOT / "docs/mechanics/DESIGN_DOCTRINE.md",
    ROOT / "docs/mechanics/GAMEPLAY_LOOP_CANON.md",
    ROOT / "docs/mechanics/GAMEPLAY_LOOP_EXAMPLES.md",
]

GOLD = "B8860B"
DARK = "1D2733"
PALE = "F7F1E3"
RED = "8B2E2E"
GRAY = "59636E"


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def configure_styles(doc: Document):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    for name, size, color, before, after in [
        ("Title", 28, DARK, 12, 12),
        ("Subtitle", 14, GRAY, 0, 12),
        ("Heading 1", 18, DARK, 16, 7),
        ("Heading 2", 14, RED, 12, 5),
        ("Heading 3", 11.5, GOLD, 8, 3),
    ]:
        style = styles[name]
        style.font.name = "Aptos Display" if name != "Normal" else "Aptos"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    styles["List Bullet"].font.name = "Aptos"
    styles["List Number"].font.name = "Aptos"


def add_header_footer(doc: Document):
    for section in doc.sections:
        section.top_margin = Inches(0.68)
        section.bottom_margin = Inches(0.68)
        section.left_margin = Inches(0.72)
        section.right_margin = Inches(0.72)
        header = section.header
        p = header.paragraphs[0]
        p.text = "DECENNIUM DESCENT  |  GAMEPLAY LOOP BIBLE"
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.name = "Aptos"
            r.font.size = Pt(8)
            r.font.color.rgb = RGBColor.from_string(GRAY)
        footer = section.footer
        p = footer.paragraphs[0]
        add_page_number(p)
        for r in p.runs:
            r.font.name = "Aptos"
            r.font.size = Pt(8)
            r.font.color.rgb = RGBColor.from_string(GRAY)


def add_run_markup(paragraph, text: str):
    # Lightweight markdown: bold, inline code, and italics.
    token_re = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)")
    pos = 0
    for match in token_re.finditer(text):
        if match.start() > pos:
            paragraph.add_run(text[pos:match.start()])
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGBColor.from_string(RED)
        else:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        pos = match.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def parse_table(lines: list[str]) -> list[list[str]]:
    rows = []
    for line in lines:
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        rows.append(cells)
    if len(rows) >= 2 and all(re.fullmatch(r":?-{3,}:?", c.replace(" ", "")) for c in rows[1]):
        rows.pop(1)
    return rows


def add_table(doc: Document, rows: list[list[str]]):
    if not rows:
        return
    cols = max(len(row) for row in rows)
    table = doc.add_table(rows=0, cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for ridx, row in enumerate(rows):
        row_obj = table.add_row()
        tr_pr = row_obj._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
        if ridx == 0:
            tbl_header = OxmlElement("w:tblHeader")
            tbl_header.set(qn("w:val"), "true")
            tr_pr.append(tbl_header)
        cells = row_obj.cells
        for cidx in range(cols):
            cell = cells[cidx]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            text = row[cidx] if cidx < len(row) else ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_run_markup(p, text)
            for run in p.runs:
                run.font.size = Pt(8)
                if ridx == 0:
                    run.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
            if ridx == 0:
                set_cell_shading(cell, DARK)
            elif ridx % 2 == 0:
                set_cell_shading(cell, "F1EADB")
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_markdown(doc: Document, text: str, first_h1_as_chapter=True):
    lines = text.splitlines()
    i = 0
    first_h1 = first_h1_as_chapter
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue
        if line.startswith("|`)" ):
            pass
        if line.startswith("|") and i + 1 < len(lines) and lines[i + 1].startswith("|"):
            block = []
            while i < len(lines) and lines[i].startswith("|"):
                block.append(lines[i])
                i += 1
            add_table(doc, parse_table(block))
            continue
        if line.startswith("# "):
            if first_h1:
                first_h1 = False
            else:
                doc.add_page_break()
            p = doc.add_paragraph(style="Heading 1")
            add_run_markup(p, line[2:].strip())
            i += 1
            continue
        if line.startswith("## "):
            p = doc.add_paragraph(style="Heading 2")
            add_run_markup(p, line[3:].strip())
            i += 1
            continue
        if line.startswith("### "):
            p = doc.add_paragraph(style="Heading 3")
            add_run_markup(p, line[4:].strip())
            i += 1
            continue
        if line.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.35)
            p.paragraph_format.right_indent = Inches(0.2)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(8)
            run = p.add_run(line[2:].strip())
            run.italic = True
            run.bold = True
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor.from_string(RED)
            i += 1
            continue
        if re.match(r"^\d+\.\s", line):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.first_line_indent = Inches(-0.22)
            add_run_markup(p, line.strip())
            i += 1
            continue
        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_run_markup(p, line[2:].strip())
            i += 1
            continue
        if line.startswith("```"):
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                code_lines.append(lines[i])
                i += 1
            i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.right_indent = Inches(0.25)
            p.paragraph_format.space_before = Pt(5)
            p.paragraph_format.space_after = Pt(7)
            run = p.add_run("\n".join(code_lines))
            run.font.name = "Consolas"
            run.font.size = Pt(8.5)
            set_paragraph_shading(p, "EEE9DF")
            continue
        # Combine adjacent ordinary lines into a paragraph.
        para_lines = [line.strip()]
        i += 1
        while i < len(lines):
            nxt = lines[i].rstrip()
            if not nxt.strip() or nxt.startswith(("#", "- ", "> ", "```", "|")) or re.match(r"^\d+\.\s", nxt):
                break
            para_lines.append(nxt.strip())
            i += 1
        p = doc.add_paragraph()
        add_run_markup(p, " ".join(para_lines))


def set_paragraph_shading(paragraph, fill: str):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def add_cover(doc: Document):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(90)
    run = p.add_run("DECENNIUM DESCENT")
    run.font.name = "Aptos Display"
    run.font.size = Pt(34)
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(DARK)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("GAMEPLAY LOOP BIBLE")
    run.font.name = "Aptos Display"
    run.font.size = Pt(26)
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(RED)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(20)
    run = p.add_run("Preparation · Expedition · Encounter · Extraction · Haven · Cycle")
    run.font.name = "Aptos"
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor.from_string(GOLD)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(75)
    run = p.add_run("Version 1.1  |  August 2026")
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor.from_string(GRAY)

    doc.add_page_break()


def add_contents(doc: Document):
    doc.add_paragraph("Contents", style="Heading 1")
    entries = [
        "Design Doctrine",
        "The Complete Gameplay Loop",
        "Modes of Play",
        "Core Resolution",
        "Expedition Organization",
        "Opening Clock and Watches",
        "Routes, Pressure, and Team Activities",
        "Supply, Strain, Integrity, Load, and Camps",
        "Encounter, Retreat, Wounds, and Encounter Construction",
        "Harvesting and Extraction",
        "Haven Resolution, Reputation, and Advancement",
        "Seasonal and Decennial Play",
        "Worked Examples",
        "Director Quick Procedures",
        "Quick Reference Tables",
    ]
    for idx, entry in enumerate(entries, 1):
        p = doc.add_paragraph(style="List Number")
        p.add_run(entry)
    p = doc.add_paragraph()
    p.add_run("This is a static contents list. Word users may insert an automatic table of contents from the document headings.").italic = True
    doc.add_page_break()


def build():
    doc = Document()
    configure_styles(doc)
    add_header_footer(doc)
    add_cover(doc)
    add_contents(doc)
    for idx, source in enumerate(SOURCES):
        if idx:
            doc.add_page_break()
        add_markdown(doc, source.read_text(encoding="utf-8"), first_h1_as_chapter=True)
    # Final appendix with playtest status.
    doc.add_page_break()
    doc.add_paragraph("Appendix: Validation Status", style="Heading 1")
    p = doc.add_paragraph()
    p.add_run("Rules status. ").bold = True
    p.add_run("The gameplay loop is complete and canonical. Cursor simulations confirmed procedural ambiguities that were corrected in version 1.1. Their model-generated combat numbers are not treated as human playtest proof. Numerical encounter budgets, feat interactions, high-level class combinations, economic throughput, long-term seasonal progression, and actual reading speed still require dice-level or human table evidence.")
    p = doc.add_paragraph()
    p.add_run("Design rule. ").bold = True
    p.add_run("Decennium Descent should have 5e's readability and speed, Pathfinder 2e's tactical clarity and balance discipline, and its own expedition, extraction, Haven, and seasonal systems.")
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
