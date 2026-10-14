"""Director-only completion of T1–5 rules package (Silent Ones is in a table cell)."""
from __future__ import annotations

import shutil
import sys
from datetime import datetime
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[2]
DIRECTOR = ROOT / "docs" / "manuals" / "The_Directors_Crucible_Complete_Director_Edition.docx"
BACKUP_DIR = ROOT / "backups"


def set_text(paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for r in paragraph.runs[1:]:
            r.text = ""
    else:
        paragraph.add_run(text)


def set_cell_text(cell, text: str) -> None:
    paras = cell.paragraphs
    if not paras:
        cell.add_paragraph(text)
        return
    set_text(paras[0], text)
    for p in paras[1:]:
        set_text(p, "")


def find_prefix(doc: Document, prefix: str):
    for p in doc.paragraphs:
        if (p.text or "").strip().startswith(prefix):
            return p
    raise SystemExit(f"MISSING prefix: {prefix!r}")


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8")
    BACKUP_DIR.mkdir(exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup = BACKUP_DIR / f"Complete_Director_Edition_pre_t1_5_rules_{stamp}.docx"
    shutil.copy2(DIRECTOR, backup)
    print("BACKUP", backup)

    doc = Document(str(DIRECTOR))
    edits: list[str] = []

    awards = find_prefix(doc, "Treasure and awards. The cumulative level 1–20 XP thresholds")
    set_text(
        awards,
        "Treasure and awards. The cumulative level 1–20 XP thresholds in the Player "
        "Rulebook are canon for level and proficiency when XP tracking is in force. "
        "For Opening campaigns (Tier 1–5 through a Decennial Opening), milestone "
        "advancement is official: award levels at Gate seals, major arc beats, or "
        "other milestones announced before the Opening begins. No published CR→XP "
        "or encounter-budget table exists in this book or the Player Rulebook. "
        "Companion Session CR→XP helpers are optional and non-canon — label them "
        "playtest if used; they do not write to character saves. Outside Opening "
        "scope, specific encounter awards remain Director judgment unless a published "
        "adventure provides them.",
    )
    edits.append("awards")

    burnout = find_prefix(doc, "Overcast and Burnout.")
    set_text(
        burnout,
        "Overcast and Burnout. Extended, Subtle, and Distant each cost +1 RP; "
        "Empowered costs +2; Widened, Quickened, and Twinned cost +3; Heightened "
        "costs +4. Burnout risk begins when the remaining pool falls below one "
        "quarter of maximum after a cast. Until a final Burnout table is published, "
        "Burnout is permanent GM fiat: announce a playtest consequence before it "
        "applies (default suggestion: disadvantage on Weave-related checks until "
        "the end of the caster's next turn) and record how it affected play. "
        "Companion Session Burnout badges are optional reminders only.",
    )
    edits.append("burnout")

    pool = find_prefix(doc, "Pool size. Use the class tables in the Player Rulebook.")
    set_text(
        pool,
        "Pool size. Use the class tables in the Player Rulebook. Wizard, Mage, "
        "Cleric, and Druid have complete level 1–20 Resonance tables. When a "
        "Weave-casting class lacks a finalized table, treat companion draft RP "
        "pool estimates as playtest material only — write the chosen progression "
        "on the character sheet and say Draft aloud at session zero.",
    )
    edits.append("pool")

    ruling = find_prefix(doc, "When a ruling interacts with a class limit.")
    set_text(
        ruling,
        "When a ruling interacts with a class limit. Let the most specific published "
        "class feature, action-economy limit, copy limit, or aura limit control for "
        "that session. Draft class ceiling caps (companion Class Ceiling Caps callouts "
        "and CLASS_CEILING_CAPS.md) are permanent GM-fiat playtest guardrails until "
        "author-promoted — announce them before they bind a player, record ambiguity, "
        "and clarify before the next session instead of relitigating mid-fight.",
    )
    edits.append("ceiling")

    always_xp = find_prefix(doc, "Label provisional combat statistics, XP awards")
    set_text(
        always_xp,
        "Label provisional combat statistics and any optional CR→XP companion helpers "
        "as playtest / non-canon. Opening campaigns use milestone advancement as "
        "official; do not present Session draft XP as rulebook awards.",
    )
    edits.append("always-xp")

    guard = find_prefix(doc, "Director-secret: the Silent Ones, later called the Ul")
    set_text(
        guard,
        "Director-secret: Silent Ones ARE the Old Ones (same entity). The Silent "
        "Ones, later called the Ul'nak, built the underlying Heartcell.",
    )
    edits.append("guardrails")

    updated_table = False
    for ti, table in enumerate(doc.tables):
        for row in table.rows:
            for cell in row.cells:
                t = cell.text.strip()
                if "Director-secret. The Silent Ones" in t and "reward architecture" in t:
                    set_cell_text(
                        cell,
                        "🔒 Director-secret. Silent Ones ARE the Old Ones (same entity). "
                        "Register note: Silent Ones is the GM-secret framing; Old Ones / Ul'nak "
                        "is the Matrix-of-Oblivion cosmology framing — not a second force. "
                        "The Silent Ones — the Old Ones called the Ul'nak after the Sundering — "
                        "built the machine beneath the public Crucible: the Heartcell. It contains "
                        "the hundred-floor resonance structure, the ten Gate divisions, the "
                        "recurring Opening, and the reward architecture that makes loot behave as "
                        "bait. Calan'thir and Asmodeus did not originate that machine. After the "
                        "Old Gods were barely defeated, they found it, claimed it, modified it, "
                        "activated it as a resonance valve, and bound the Divine Pact to its "
                        "operation. Both histories are therefore true at different layers: the "
                        "Silent Ones built the prison; Calan'thir and Asmodeus built the Crucible "
                        "the world enters.",
                    )
                    edits.append(f"table{ti}-silent-ones")
                    updated_table = True
                if t == "Director-secret builders associated with the Old Ones and the Matrix":
                    set_cell_text(
                        cell,
                        "Director-secret builders; Silent Ones ARE the Old Ones (same entity). "
                        "Ul'nak is the post-Sundering / Matrix framing of that same force.",
                    )
                    edits.append(f"table{ti}-gloss")

    if not updated_table:
        raise SystemExit("Failed to update Ch 8 Silent Ones table cell")

    doc.save(str(DIRECTOR))
    print("SAVED", DIRECTOR)
    print("EDITS", edits)


if __name__ == "__main__":
    main()
