# -*- coding: utf-8 -*-
"""Fix DMG Session Zero public-myth mix-up. Backup first. No commit."""
from __future__ import annotations

import shutil
from datetime import datetime
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[2]
SRC = ROOT / "docs/manuals/The_Directors_Crucible_Complete_Director_Edition.docx"
BAK_DIR = ROOT / "backups"


def main() -> None:
    BAK_DIR.mkdir(exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    bak = BAK_DIR / f"The_Directors_Crucible_Complete_Director_Edition_pre_continuity_audit_{stamp}.docx"
    shutil.copy2(SRC, bak)
    print(f"backup: {bak}")

    new = (
        "2. Explain the public Crucible myth. Calan\u2019thir and Asmodeus forged the Divine Pact "
        "and jointly built the Crucible known to the world. That joint-creation story is the complete "
        "public explanation at campaign start. Do not mention the Heartcell, Silent Ones, Ul\u2019nak, "
        "or Matrix of Oblivion (you hold the secret discoverer/activator layer separately \u2014 Chapter 8)."
    )

    doc = Document(str(SRC))
    hits = 0

    def patch_paragraph(p) -> bool:
        t = p.text
        if "public Crucible myth" in t and "discovered, claimed, modified" in t:
            print("FOUND:", repr(t[:200]))
            if p.runs:
                p.runs[0].text = new
                for r in p.runs[1:]:
                    r.text = ""
            else:
                p.text = new
            return True
        return False

    for p in doc.paragraphs:
        if patch_paragraph(p):
            hits += 1

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if patch_paragraph(p):
                        hits += 1

    print("hits:", hits)
    if hits == 0:
        print("NO HIT — listing candidates:")
        for p in doc.paragraphs:
            if "public Crucible myth" in p.text or "discovered, claimed" in p.text:
                print(repr(p.text[:300]))
        return

    doc.save(str(SRC))
    print("saved:", SRC)


if __name__ == "__main__":
    main()
