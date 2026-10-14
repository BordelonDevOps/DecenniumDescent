#!/usr/bin/env python3
"""Final pass: find remaining live wrong Opening 240 and unmarked 120-as-PoNR."""
from __future__ import annotations

import re
import sys
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
ROOT = Path(r"C:\Github\Descent assets\DecenniumDescent-main")

SKIP = {"node_modules", ".git", ".next", "dist", "build", "backups", "__pycache__"}
EXTS = {".md", ".txt", ".json", ".docx"}

try:
    from docx import Document
except ImportError:
    Document = None


def read(path: Path) -> str:
    if path.suffix.lower() == ".docx":
        doc = Document(str(path))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for t in doc.tables:
            for row in t.rows:
                for c in row.cells:
                    if c.text.strip():
                        parts.append(c.text)
        return "\n".join(parts)
    raw = path.read_bytes()
    for enc in ("utf-8", "utf-16", "utf-16-le", "latin-1"):
        try:
            return raw.decode(enc)
        except Exception:
            continue
    return ""


# Wrong Opening as live canon (not 2400)
OPENING_FAIL = re.compile(
    r"(?:"
    r"\b240[\-\s]?hour(?:s)?\s+[Oo]pening\b|"
    r"[Oo]pening\s+(?:lasts|is|of|=|:)\s*240\s*hours?\b|"
    r"[Oo]pening.{0,30}\b240\s*h(?:ours?)?\b(?!ours?\s*/|\s*/\s*100)|"
    r"ten[\-\s]?day\s+[Oo]pening|"
    r"10[\-\s]?day\s+[Oo]pening|"
    r"do not multiply the tabletop clock by ten|"
    r"The 240-Hour Descent\b(?!\s*\(100)|"
    r"have 240 hours\b|"
    r"blank 240-hour|"
    r"full 240-hour Opening|"
    r"seal closes at hour 240\b|"
    r"measured as 240 hours"
    r")",
    re.I,
)

# Horn at 120 as PoNR without debt/typo label nearby
HORN_PONR_120 = re.compile(
    r"(?:"
    r"At hour 120,\s*day 5|"
    r"Horn sounds at hour 120\b|"
    r"Horn at hour 120\b|"
    r"with the Horn at hour 120\b|"
    r"giant horn sounds at 120 hours|"
    r"horn sounds at 120 hours|"
    r"At 120 hours underground, a giant horn|"
    r"At 120 hours, a giant horn|"
    r"Horn Clock:\s*at 120 hours|"
    r"horn clock \(120h|"
    r"120h horn|"
    r"120 h horn|"
    r"120-hour underground horn|"
    r"HORN_AT\s*=\s*120"
    r")",
    re.I,
)

OK = re.compile(
    r"typo|superseded|drift|wrongly|missing zero|amended|code debt|code/doc debt|"
    r"was\s+\*?\*?\s*240|still reads|do not change|pending|ask owner|historical|"
    r"intentionally left|not changed|flag(?:ged)?|may be wrong|"
    r"owner lock round 2|round-2|Q-OV-1|Q4|"
    r"240\s*[→\-–>]+\s*2400|compress the Opening to 240|"
    r"prior 120|legacy|record of|sim(?:ulation)?|play report|ended \d+/120|"
    r"/\s*120 h\)|banked|buffer before 120|≤120|<=120|"
    r"soft quit|Canticle|Floor 50",
    re.I,
)

# Negative look: hour 1200 should not match hour 120
def has_hour_120_not_1200(text: str) -> bool:
    # replace 1200 and 2400 temporarily
    t = re.sub(r"1200|2400", "XXXX", text)
    return bool(HORN_PONR_120.search(t)) or bool(re.search(r"hour\s*120\b|120\s*h(?:ours?)?\b", t, re.I) and re.search(r"horn|ponr|opening", t, re.I))


results = {"OPENING_FAIL": [], "HORN_LIVE_FAIL": [], "HORN_OK_OR_DEBT": [], "SOFT_QUIT": []}

for dirpath, dirnames, filenames in __import__("os").walk(ROOT):
    dirnames[:] = [d for d in dirnames if d not in SKIP and not d.startswith(".")]
    for fn in filenames:
        path = Path(dirpath) / fn
        if path.suffix.lower() not in EXTS:
            continue
        if re.search(r"\.pre-|pre-hybridb|_scan_|_analyze_|_target_|_opening_clock_scan|OPENING_CLOCK_CONTINUITY", path.name):
            continue
        try:
            text = read(path)
        except Exception as e:
            continue
        rel = str(path.relative_to(ROOT)).replace("\\", "/")
        for i, line in enumerate(text.splitlines(), 1):
            if "Canticle" in line or "soft quit" in line.lower() or "Floor 50" in line and "horn" in line.lower():
                if re.search(r"horn|ponr|opening|canticle|soft quit", line, re.I):
                    results["SOFT_QUIT"].append((rel, i, line.strip()[:200]))
            # Opening fail
            if OPENING_FAIL.search(line):
                # avoid 2400 false positive
                tmp = re.sub(r"2400", "XXXX", line)
                if not OPENING_FAIL.search(tmp):
                    continue
                if OK.search(line):
                    results["HORN_OK_OR_DEBT"].append((rel, i, "opening-okctx", line.strip()[:200]))
                else:
                    results["OPENING_FAIL"].append((rel, i, line.strip()[:240]))
            # Horn 120
            tmp = re.sub(r"1200|2400", "XXXX", line)
            if HORN_PONR_120.search(tmp) or (
                re.search(r"(?:horn|ponr).{0,40}120\s*h|120\s*h.{0,40}(?:horn|ponr)|giant horn sounds at 120", tmp, re.I)
            ):
                if OK.search(line):
                    results["HORN_OK_OR_DEBT"].append((rel, i, "horn-okctx", line.strip()[:200]))
                else:
                    # playtest historical with numbers like 106/120 often OK
                    if re.search(r"\d+\s*/\s*120|sim|play report|T4_|SIM_", rel + line, re.I) and "hornRule" not in line:
                        results["HORN_OK_OR_DEBT"].append((rel, i, "hist", line.strip()[:200]))
                    else:
                        results["HORN_LIVE_FAIL"].append((rel, i, line.strip()[:240]))

print("OPENING_FAIL", len(results["OPENING_FAIL"]))
for r in results["OPENING_FAIL"][:40]:
    print(f"  {r[0]}:{r[1]}: {r[2]}")

print("\nHORN_LIVE_FAIL", len(results["HORN_LIVE_FAIL"]))
by = {}
for r in results["HORN_LIVE_FAIL"]:
    by.setdefault(r[0], []).append(r)
for f, items in sorted(by.items(), key=lambda x: -len(x[1]))[:35]:
    print(f"=== {f} ({len(items)})")
    for it in items[:6]:
        print(f"  L{it[1]}: {it[2]}")

print("\nSOFT_QUIT sample", len(results["SOFT_QUIT"]))
for r in results["SOFT_QUIT"][:20]:
    print(f"  {r[0]}:{r[1]}: {r[2]}")
