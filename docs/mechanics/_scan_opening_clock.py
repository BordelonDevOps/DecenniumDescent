#!/usr/bin/env python3
"""Scan Opening/Horn clock references for continuity check 2026-07-31."""
from __future__ import annotations

import json
import os
import re
from pathlib import Path

ROOT = Path(r"C:\Github\Descent assets\DecenniumDescent-main")
SKIP_DIRS = {
    "node_modules",
    ".git",
    ".next",
    "dist",
    "build",
    "coverage",
    "__pycache__",
}
TEXT_EXTS = {
    ".md",
    ".txt",
    ".json",
    ".ts",
    ".tsx",
    ".js",
    ".jsx",
    ".mjs",
    ".cjs",
}

RELEVANCE = re.compile(
    r"opening|horn|ponr|descent[\s\-]?window|2400|1200|240\s*h|120\s*h|"
    r"240\s*hour|120\s*hour|100[\-\s]?day|ten[\-\s]?day|10[\-\s]?day|"
    r"multiply by ten|multiply by 10|point of no return|canticle|soft quit|"
    r"hour\s*120|hour\s*240|day\s*50|day\s*5\b|HORN_AT|hornRule|"
    r"\b(?:2400|1200|240|120)\s*h(?:ours?)?\b|"
    r"(?:hour|at)\s*(?:2400|1200|240|120)\b",
    re.I,
)

TAG_PATTERNS = [
    (re.compile(r"2400\s*h(?:ours?)?", re.I), "correct_2400h"),
    (re.compile(r"\b2400\b"), "num_2400"),
    (re.compile(r"100[\-\s]?day(?:s)?\s+Opening|Opening.{0,40}100\s*days?", re.I), "correct_100day"),
    (re.compile(r"hour\s*1200|1200\s*h(?:ours?)?|\b1200\b", re.I), "num_1200"),
    (re.compile(r"day\s*50|Day\s*50", re.I), "day_50"),
    (re.compile(r"240\s*h(?:ours?)?", re.I), "suspect_240h"),
    (re.compile(r"\b240\b"), "num_240"),
    (re.compile(r"ten[\-\s]?day\s+Opening|10[\-\s]?day\s+Opening", re.I), "suspect_10day"),
    (re.compile(r"hour\s*120(?!\d)|120\s*h(?:ours?)?", re.I), "suspect_120h"),
    (re.compile(r"\b120\b"), "num_120"),
    (re.compile(r"day\s*5\b|Day\s*5\b", re.I), "day_5"),
    (re.compile(r"do not multiply by ten|don.?t multiply by ten|not multiply by 10", re.I), "multiply_defense"),
    (re.compile(r"Opening", re.I), "opening_word"),
    (re.compile(r"Horn|PoNR|point of no return|descent[\-\s]?window", re.I), "horn_window"),
    (re.compile(r"HORN_AT|hornRule", re.I), "code_horn"),
]


def extract_docx(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    parts: list[str] = []
    for p in doc.paragraphs:
        if p.text.strip():
            parts.append(p.text)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)
    return "\n".join(parts)


def read_text(path: Path) -> str | None:
    raw = path.read_bytes()
    for enc in ("utf-8", "utf-16", "utf-16-le", "latin-1"):
        try:
            return raw.decode(enc)
        except Exception:
            continue
    return None


def classify_line(text: str, tags: list[str], file: str) -> str:
    t = text.lower()
    # Explicit OK-CONTEXT markers
    ok_markers = [
        "typo",
        "superseded",
        "drift",
        "code debt",
        "code/doc debt",
        "wrongly stated",
        "missing zero",
        "owner-confirmed typo",
        "amended to",
        "was a typo",
        "historical",
        "changelog",
        "do not change",
        "pending",
        "ask owner",
        "q-ov-1",
        "round 1",
        "~~",
        "struck",
        "incorrect",
        "must not remain",
        "rejected",
        "false positive",
    ]
    if any(m in t for m in ok_markers) and (
        "240" in t or "120" in t or "typo" in t or "debt" in t or "superseded" in t
    ):
        # Still FAIL if it asserts live wrong canon alongside callout? Prefer OK-CONTEXT for lock docs.
        if re.search(r"is canon.*\b240\b|\b240\b.*is canon|opening.{0,30}\b240\b(?!0)", t) and "typo" not in t and "superseded" not in t and "wrong" not in t:
            pass
        else:
            return "OK-CONTEXT"

    # Correct patterns
    has_2400 = "correct_2400h" in tags or ("num_2400" in tags and ("opening" in t or "2400" in t))
    has_1200 = "num_1200" in tags and ("horn" in t or "ponr" in t or "1200" in t or "day 50" in t or "day50" in t)
    has_100 = "correct_100day" in tags or ("100 day" in t or "100-day" in t)

    # False positive filters for bare 120/240
    false_pos = False
    if "num_120" in tags or "suspect_120h" in tags:
        if re.search(
            r"120\s*(?:feet|ft\.?|xp|ac)|cr\s*12|floor\s*12|chapter\s*12|"
            r"expedition.{0,20}12|12[\-\s]?person|12\s*characters|"
            r"6[\–\-]12\s*hour|12\s*hour session|max(?:imum)?\s*12|"
            r"appendix\s*b|level\s*12|\+12\b|dc\s*12",
            t,
            re.I,
        ):
            false_pos = True
        # 120 without hour/horn context
        if "suspect_120h" not in tags and "horn" not in t and "opening" not in t and "ponr" not in t:
            if not re.search(r"120\s*h|hour\s*120\b", t, re.I):
                false_pos = True

    if "num_240" in tags or "suspect_240h" in tags:
        if re.search(r"240\s*(?:feet|ft\.?|xp)|cr\s*24|floor\s*24|chapter\s*24", t, re.I):
            false_pos = True
        if "suspect_240h" not in tags and "opening" not in t and "descent" not in t and "horn" not in t:
            if not re.search(r"240\s*h|hour\s*240\b", t, re.I):
                false_pos = True

    # FAIL: live wrong canon
    live_wrong_240 = bool(
        re.search(
            r"(?:opening|descent).{0,60}(?:240\s*h|240\s*hours?|\b240\b(?!0))|"
            r"(?:240\s*h|240\s*hours?).{0,60}(?:opening|descent)|"
            r"the\s+240[\-\s]?hour\s+descent|"
            r"ten[\-\s]?day\s+opening|10[\-\s]?day\s+opening",
            t,
            re.I,
        )
    )
    # Exclude if clearly saying 2400
    if "2400" in t and re.search(r"240(?!0)", t):
        # mixed - check if 240 is called typo
        if any(m in t for m in ("typo", "wrong", "superseded", "drift", "was", "not", "≠", "!=")):
            live_wrong_240 = False
        elif re.search(r"\b240\b(?!0)", t) and "2400" in t:
            # e.g. "240 → 2400" is OK-CONTEXT
            if re.search(r"240\s*[→\-–].*2400|2400.*(?:from|was|not)\s*240|240\s*hours?.{0,20}typo", t):
                live_wrong_240 = False

    live_wrong_120 = bool(
        re.search(
            r"(?:horn|ponr|point of no return).{0,50}(?:120\s*h|120\s*hours?|hour\s*120\b|day\s*5\b)|"
            r"(?:120\s*h|120\s*hours?|hour\s*120\b).{0,50}(?:horn|ponr|point of no return)|"
            r"horn.{0,20}(?:at|=|:)\s*120\b|"
            r"HORN_AT\s*=\s*120",
            t,
            re.I,
        )
    )
    if "1200" in t and live_wrong_120:
        if any(m in t for m in ("typo", "wrong", "superseded", "debt", "may be", "pending", "do not change")):
            live_wrong_120 = False
        elif re.search(r"120\s*[→\-–].*1200|1200.*(?:from|was|not)\s*120", t):
            live_wrong_120 = False

    multiply = "multiply_defense" in tags

    # Code paths: implementation debt separate
    is_code = file.replace("\\", "/").startswith("src/") or file.endswith(
        (".ts", ".tsx", ".js", ".jsx")
    )

    if multiply and ("240" in t or "opening" in t):
        if any(m in t for m in ("typo", "wrong", "do not", "must not", "reject")):
            return "OK-CONTEXT"
        return "FAIL"

    if false_pos and not live_wrong_240 and not live_wrong_120:
        return "OK-CONTEXT"  # treat filtered as non-issue; won't count as FAIL

    if is_code and (live_wrong_120 or "HORN_AT" in text or "hornRule" in text):
        if re.search(r"120", t) and ("horn" in t or "HORN_AT" in text or "hornRule" in text):
            return "CODE-DEBT"

    if live_wrong_240 or live_wrong_120:
        # docs that are continuity stamps documenting the error
        if any(
            x in file.replace("\\", "/")
            for x in (
                "CONTINUITY_SOURCE_OF_TRUTH",
                "DESCENT_WINDOW_CANON",
                "OWNER_COSMOLOGY",
                "DIRECTOR_BIBLE_CLOCK_TYPO",
                "OPENING_CLOCK_CONTINUITY",
                "_scan_opening",
                "_opening_clock_scan",
            )
        ):
            return "OK-CONTEXT"
        if any(m in t for m in ("typo", "superseded", "debt", "wrongly", "missing zero", "drift")):
            return "OK-CONTEXT"
        return "FAIL"

    if has_2400 or has_1200 or has_100 or "day_50" in tags:
        if ("suspect_240h" in tags or "suspect_10day" in tags) and "2400" not in t:
            return "FAIL"
        if "suspect_120h" in tags and "1200" not in t and "horn" in t:
            return "AMBIGUOUS"
        return "PASS"

    if "suspect_240h" in tags or "suspect_10day" in tags:
        return "FAIL"
    if "suspect_120h" in tags and ("horn" in t or "opening" in t or "ponr" in t):
        return "AMBIGUOUS"
    if "opening_word" in tags or "horn_window" in tags:
        return "AMBIGUOUS"
    return "AMBIGUOUS"


def main() -> None:
    hits: list[dict] = []
    errors: list[tuple[str, str]] = []
    files_scanned = 0

    for dirpath, dirnames, filenames in os.walk(ROOT):
        dirnames[:] = [d for d in dirnames if d not in SKIP_DIRS and not d.startswith(".")]
        for fn in filenames:
            path = Path(dirpath) / fn
            # skip our own scan artifacts / this script output mid-run
            if path.name.startswith("_scan_opening") or path.name.startswith("_opening_clock_scan"):
                continue
            if path.name == "OPENING_CLOCK_CONTINUITY_CHECK_2026-07-31.md":
                continue
            ext = path.suffix.lower()
            try:
                if ext in TEXT_EXTS:
                    text = read_text(path)
                    if text is None:
                        continue
                elif ext == ".docx":
                    try:
                        text = extract_docx(path)
                    except Exception as e:
                        errors.append((str(path.relative_to(ROOT)), f"docx:{e}"))
                        continue
                else:
                    continue
            except Exception as e:
                errors.append((str(path.relative_to(ROOT)), str(e)))
                continue

            files_scanned += 1
            rel = str(path.relative_to(ROOT)).replace("\\", "/")
            for i, line in enumerate(text.splitlines(), 1):
                if not RELEVANCE.search(line):
                    continue
                clipped = line.strip()
                if len(clipped) > 500:
                    clipped = clipped[:500] + "…"
                tags = [name for pat, name in TAG_PATTERNS if pat.search(line)]
                cls = classify_line(clipped, tags, rel)
                # Drop pure false-positive OK-CONTEXT with no clock interest
                if cls == "OK-CONTEXT" and not any(
                    x in tags
                    for x in (
                        "suspect_240h",
                        "suspect_120h",
                        "suspect_10day",
                        "num_240",
                        "num_120",
                        "multiply_defense",
                        "correct_2400h",
                        "num_2400",
                        "num_1200",
                        "day_50",
                        "code_horn",
                    )
                ):
                    # keep if mentions wrong clocks in continuity docs
                    if not re.search(r"\b240\b|\b120\b|typo|debt|superseded", clipped, re.I):
                        continue
                hits.append(
                    {
                        "file": rel,
                        "line": i,
                        "text": clipped,
                        "tags": tags,
                        "class": cls,
                    }
                )

    out = ROOT / "docs" / "mechanics" / "_opening_clock_scan_raw.json"
    out.write_text(
        json.dumps({"files_scanned": files_scanned, "hits": hits, "errors": errors}, indent=2),
        encoding="utf-8",
    )

    from collections import Counter, defaultdict

    counts = Counter(h["class"] for h in hits)
    by_class: dict[str, list] = defaultdict(list)
    for h in hits:
        by_class[h["class"]].append(h)

    print(f"FILES_SCANNED={files_scanned}")
    print(f"HITS={len(hits)}")
    print(f"ERRORS={len(errors)}")
    print("COUNTS", dict(counts))
    for cls in ("FAIL", "AMBIGUOUS", "CODE-DEBT", "OK-CONTEXT", "PASS"):
        items = by_class.get(cls, [])
        print(f"\n===== {cls} ({len(items)}) =====")
        # group by file
        byf: dict[str, list] = defaultdict(list)
        for h in items:
            byf[h["file"]].append(h)
        for f, fl in sorted(byf.items(), key=lambda x: -len(x[1]))[:50]:
            print(f"--- {f} ({len(fl)})")
            for it in fl[:8]:
                print(f"  L{it['line']}: {it['text'][:180]}")


if __name__ == "__main__":
    main()
