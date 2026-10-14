"""Surgical Player Rulebook alignment to Director's Crucible (player-facing only)."""
from __future__ import annotations

import shutil
import sys
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parents[2]
SRC = ROOT / "decennium descent player rulebook.docx"
BACKUP_DIR = ROOT / "backups"


def set_text(paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for r in paragraph.runs[1:]:
            r.text = ""
    else:
        paragraph.add_run(text)


def insert_after(paragraph, text: str) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    new_para.add_run(text)
    return new_para


def index_paras(doc: Document) -> dict[str, Paragraph]:
    out: dict[str, Paragraph] = {}
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if t and t not in out:
            out[t] = p
    return out


def must(by_text: dict[str, Paragraph], *, exact: str | None = None, prefix: str | None = None) -> Paragraph:
    if exact is not None:
        if exact not in by_text:
            raise SystemExit(f"MISSING exact: {exact[:100]!r}")
        return by_text[exact]
    assert prefix is not None
    for t, p in by_text.items():
        if t.startswith(prefix):
            return p
    raise SystemExit(f"MISSING prefix: {prefix!r}")


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8")
    BACKUP_DIR.mkdir(exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup = BACKUP_DIR / f"decennium_descent_player_rulebook_pre_dmg_align_{stamp}.docx"
    shutil.copy2(SRC, backup)
    print("BACKUP", backup)

    doc = Document(str(SRC))
    edits: list[str] = []
    by_text = index_paras(doc)

    # --- Opening clocks + public Crucible (before Expeditions) ---
    paras = list(doc.paragraphs)
    i_exp = next(
        i for i, p in enumerate(paras) if (p.text or "").strip() == "Expeditions in Aethel"
    )
    i_prev = i_exp - 1
    while i_prev > 0 and not (paras[i_prev].text or "").strip():
        i_prev -= 1
    anchor = paras[i_prev]

    blocks = [
        "The Decennial Opening",
        (
            "Once each decade, the Shardspire activates and the Decennium — the hundred-floor Great Descent — "
            "accepts mortal Expeditions again. Between Cycles the dungeon is sealed. Survivors, maps, and rumor "
            "carry knowledge forward until the next Opening."
        ),
        (
            "What every character in Aethel is taught about the Crucible is not a trick: after the Sundering, "
            "Calan'thir and Asmodeus forged the Divine Pact and jointly built the Crucible — also called "
            "Calan'thir's Trial — as neutral proving ground and a monument to Order and Hell managing their "
            "differences instead of ending them by force. Cultures argue which sovereign deserves more credit. "
            "Virtually no one disputes that both built it together. This is the complete public explanation "
            "available at the start of a campaign."
        ),
        (
            "The Opening lasts 100 days — 2400 hours — from Shardspire activation. Track it as a shared calendar "
            "your Expedition lives inside, not only as a private countdown."
        ),
        (
            "The Horn sounds at hour 1200, day 50, the exact midpoint of the Opening. After the Horn, known "
            "extraction routes become less reliable, rescue contracts thin out, and allies begin making final "
            "decisions about whether to push deeper or turn upward. The Horn does not seal the dungeon by itself, "
            "but it ends the illusion that every Expedition can still recover from one more mistake."
        ),
        (
            "At hour 2400, day 100, the Decennium seals. Anyone still inside is not automatically killed; survivors "
            "may be trapped until the next Cycle, ten years later. Surface cultures and Expeditions treat the "
            "closing as an expected, telegraphed end of the window — not a surprise ambush by the rules."
        ),
        (
            "Many Expeditions that reach the Sanctum of Echoing Choirs (Tier 5, floors 41–50) turn back there "
            "because gear, nerve, and Opening hours have already been spent. The Bar of Last Refuge is the "
            "cultural name for that threshold: a place where parties decide whether the next push is worth the "
            "risk. Soft turn-back at this depth is not the same event as the Horn."
        ),
        (
            "Each tier of the Descent ends in a Gate sealed by a Gate Guardian — a named seal-creature whose "
            "defeat or bypass opens the way deeper. Your Director stages Guardians when your Expedition reaches "
            "them; their full roster and tactics belong to the Director's materials."
        ),
    ]

    cur = anchor
    for text in blocks:
        cur = insert_after(cur, text)
        edits.append(f"Inserted Opening: {text[:50]}...")

    by_text = index_paras(doc)

    # --- Expedition coalition clarification ---
    exp_cap = (
        "In Decennium Descent, adventuring parties are known as Expeditions and can consist of up to 12 characters. "
        "This limit represents the total number of individuals working together toward common goals, not just those "
        "engaged in combat. An effective Expedition distributes its members across various activities to maximize "
        "efficiency, resource gathering, information collection, and overall success."
    )
    p = must(by_text, exact=exp_cap)
    clar = (
        " A force larger than twelve is not one oversized Expedition: it is two or more separate Expeditions "
        "coordinating, each still respecting the twelve-character cap."
    )
    if "two or more separate Expeditions" not in p.text:
        set_text(p, p.text.rstrip() + clar)
        edits.append("Added Expedition coalition clarification")

    by_text = index_paras(doc)

    # --- Appendix B ---
    p = must(by_text, exact="# Appendix B: The Lore of Aethel")
    set_text(p, "# Appendix B: Faiths and Cosmologies of Aethel")
    edits.append("Renamed Appendix B")

    disclaimer = (
        "The beliefs recorded in this appendix are mortal theology, not objective cosmic history. "
        "Priests, scholars, and cults across Aethel have spent millennia reasoning about forces they cannot fully "
        "perceive. What follows is what the faithful believe: a creation myth, a pantheon of worshiped faces, and "
        "a diagram of the planes — sincerely held, widely taught, and useful for understanding how ordinary people "
        "in Aethel think about the divine. Your character may treat this as literal truth. Deeper mysteries of the "
        "Spiral and the Crucible may surface only through play; even then, the wider world does not rewrite its "
        "temples overnight. Scholars may also speak of the First Choir — seven named virtues (Korvelan, Vaerilune, "
        "Delvaran, Shaelara, Xalarex, Thalirion, and Virellian) — as part of the wider Spiral story. That knowledge "
        "is not secret, but it is not tavern talk either."
    )
    insert_after(p, disclaimer)
    edits.append("Inserted Appendix B disclaimer")

    by_text = index_paras(doc)

    p = must(by_text, exact="## The Creation Myth: The Weaving of Existence")
    set_text(p, "## The Creation Myth: The Weaving of Existence (Mortal Teaching)")
    edits.append("Reframed Creation Myth heading")

    p = must(by_text, prefix="In the beginning, there was only the Void")
    set_text(
        p,
        "In the beginning, mortal theologians say, there was only the Void — a vast, formless expanse of potential. "
        "Within this nothingness they name two primordial principles: Lumina, the Radiant Consciousness, and Umbra, "
        "the Abyssal Depth. Neither light nor dark, neither good nor evil, these twin abstractions of creation are "
        "how priests reason about first causes — not laboratory descriptions of beings scholars can point to in the "
        "sky. In the folk story, for eons beyond counting Lumina and Umbra drifted through the Void, unaware of "
        "their own existence. Then, in a moment that transcends time itself, they became conscious. Lumina's first "
        "thought was of wonder; Umbra's was of curiosity. As their awareness grew, so did their desire to understand "
        "the nature of their own being.",
    )
    edits.append("Reframed Lumina/Umbra")

    p = must(
        by_text,
        prefix="As the planes stabilized following the Sundering, fragments of Lumina and Umbra",
    )
    set_text(
        p,
        "As the planes stabilized following the Sundering, folk religion teaches that fragments of Lumina and Umbra's "
        "consciousness crystallized into the first gods. These divine faces — worship masks mortals place on real "
        "cosmic forces they cannot otherwise name — embody specific aspects of those primordial principles, each "
        "representing a particular facet of existence as temples understand it.",
    )
    edits.append("Reframed Birth of the Gods")

    p = must(by_text, prefix="These fifteen Greater Gods, known collectively as the Primeval Pantheon")
    set_text(
        p,
        "These fifteen Greater Gods, known collectively as the Fifteenfold Pantheon (or Primeval Pantheon in older "
        "temple texts), shape how mortals name the fundamental nature of the cosmos. Later, as civilizations arose, "
        "Lesser Gods and cultural pantheons emerged from belief and need — the most local tier of this worship "
        "hierarchy. Most ordinary people treat the Fifteenfold as literally true; this appendix records that faith.",
    )
    edits.append("Reframed Fifteenfold naming")

    p = must(by_text, prefix="The planar structure resembles a great wheel")
    set_text(
        p,
        "Mortal cosmographers diagram the planes as a great wheel, with the Material Plane at its center and other "
        "planes arranged around it in concentric rings. This Great Wheel is a teaching tool for organizing planar "
        "reports and temple doctrine — not a claim that explorers have surveyed the complete architecture of "
        "existence. Some planes overlap with the Material Plane in places, creating regions where their influence is "
        "stronger and their denizens might cross over more easily.",
    )
    edits.append("Reframed Great Wheel")

    p = must(
        by_text,
        exact="The gods of Aethel are arranged in a complex hierarchy based on power, domain, and origin:",
    )
    set_text(
        p,
        "Temple doctrine arranges the gods of Aethel in a hierarchy based on power, domain, and origin. The list "
        "below is how mortal faith orders worship — not a Director's absolute chart of cosmic seniority:",
    )
    edits.append("Reframed Divine Hierarchy intro")

    p = must(by_text, prefix="At the top of the divine order stand the fifteen Greater Gods")
    set_text(
        p,
        "At the top of mortal divine order stand the fifteen Greater Gods of the Fifteenfold Pantheon — the worship "
        "masks folk religion places over deeper forces. Temples teach that these deities embody fundamental cosmic "
        "principles and wield power beyond mortal comprehension.",
    )
    edits.append("Reframed Primeval Pantheon blurb")

    p = must(
        by_text,
        prefix="The most ancient and devastating conflict in the cosmos is the Blood War",
    )
    set_text(
        p,
        "Temple historians call the Blood War the oldest open war among the Lower Planes: the ordered legions of "
        "Asmodeus's Nine Hells against the demonic hordes of the Abyss. In the folk telling, the conflict sharpened "
        "after the Sundering and the forging of the Divine Pact, when Asmodeus's realm of contracts and hierarchy "
        "and the Abyss's hunger for unbound ruin could no longer share a quiet border. Every escalation strains the "
        "Divine Pact that keeps Heaven and Hell from replacing mortal choice with open divine war — priests disagree "
        "whether the Blood War protects the Pact by exhausting the fiends or endangers it by feeding the very "
        "dissonance the Pact was meant to manage.",
    )
    edits.append("Rewrote Blood War lead")

    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if t.startswith("The conflict stems from fundamental philosophical differences"):
            set_text(p, "Folk doctrine summarizes the quarrel this way:")
            edits.append("Rewrote Blood War quarrel intro")
        elif t.startswith("- **Devils** believe in strict hierarchy"):
            set_text(
                p,
                "- **Devils of the Nine Hells** serve Asmodeus's law of hierarchy, contracts, and the corruption of "
                "souls through temptation and consequence",
            )
        elif t.startswith("- **Demons** embody destruction"):
            set_text(
                p,
                "- **Demons of the Abyss** embody unbound destruction, madness, and the annihilation of imposed order",
            )
        elif t.startswith("Both sides seek to expand their influence"):
            set_text(
                p,
                "Both sides press their influence across the planes, and the Material Plane often becomes a proxy "
                "battlefield. Mortals may be recruited, tricked, or coerced into serving one side or the other — "
                "usually to their eventual regret. Calan'thir's temples stress melody and restraint; Asmodeus's agents "
                "stress victory and edge. Ordinary worshipers rarely hear more than that rivalry.",
            )
            edits.append("Rewrote Blood War both-sides")
        elif t.startswith("The Blood War serves an inadvertent cosmic purpose"):
            set_text(
                p,
                "Many priests teach that the Blood War serves an inadvertent cosmic purpose: by keeping fiendish "
                "forces occupied with each other, it prevents them from uniting against the rest of the multiverse. "
                "Others whisper that the Pact itself needs the war's friction — that absolute peace would be more "
                "dangerous than endless feud. None of the Greater Gods admit either stance in public liturgy.",
            )
            edits.append("Rewrote Blood War purpose")

    by_text = index_paras(doc)
    p = must(by_text, exact="## Conclusion: The Tapestry Continues")
    set_text(p, "## Conclusion: The Tapestry Continues (As Mortals Tell It)")
    edits.append("Reframed conclusion heading")

    p = must(by_text, prefix="The world of Aethel stands at a crossroads")
    set_text(
        p,
        "In the telling of Aethel's faithful, the world stands at a crossroads. The barriers between planes feel "
        "thinner, ancient powers stir in rumor and omen, and mortal civilizations face challenges that echo past "
        "catastrophes. Yet there is also hope — new discoveries, alliances, and heroes emerging to shape the future.",
    )
    edits.append("Softened conclusion opening")

    doc.save(str(SRC))
    print("SAVED", SRC)
    print("EDIT_COUNT", len(edits))
    for e in edits:
        print("-", e)
    print("BACKUP_PATH", backup)


if __name__ == "__main__":
    main()
