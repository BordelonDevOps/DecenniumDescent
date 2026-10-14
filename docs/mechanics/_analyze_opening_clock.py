#!/usr/bin/env python3
"""Analyze opening clock scan; verify live manuals/docx."""
from __future__ import annotations

import json
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

ROOT = Path(r"C:\Github\Descent assets\DecenniumDescent-main")
data = json.loads((ROOT / "docs/mechanics/_opening_clock_scan_raw.json").read_text(encoding="utf-8"))
hits = data["hits"]

EXCLUDE_PREFIXES = ("backups/",)


def is_excluded(f: str) -> bool:
    f = f.replace("\\", "/")
    if f.startswith(EXCLUDE_PREFIXES):
        return True
    # player rulebook backups
    if re.search(r"\.pre-|pre-hybridb|pre-geography|pre-names|pre-umbrin", f, re.I):
        return True
    return False


live = [h for h in hits if not is_excluded(h["file"])]
print("LIVE_HITS", len(live), "EXCLUDED", len(hits) - len(live))
print("LIVE_COUNTS", dict(Counter(h["class"] for h in live)))

STRONG_RE = re.compile(
    r"240[\-\s]?hour\s+opening|opening.{0,40}240\s*h(?:ours?)?\b|240\s*hours?.{0,40}opening|"
    r"ten[\-\s]?day\s+opening|10[\-\s]?day\s+opening|"
    r"horn.{0,40}hour\s*120\b|hour\s*120.{0,40}horn|horn.{0,20}(?:at|=)\s*120\b|"
    r"do not multiply the tabletop|240-hour Opening|seal closes at hour 240|"
    r"blank 240-hour|full 240-hour Opening|day 5.{0,30}horn|horn.{0,30}day 5|"
    r"Opening lasts 240|measured as 240 hours|Opening is ten days",
    re.I,
)
OK_MARK = ("typo", "superseded", "amended", "was a", "wrongly", "drift", "missing zero", "code debt", "→ 2400", "-> 2400", "was **240", "still reads")

strong_fail = []
weak_fail = []
for h in live:
    if h["class"] != "FAIL":
        continue
    t = h["text"].lower()
    strong = bool(STRONG_RE.search(h["text"]))
    if any(m in t for m in OK_MARK):
        strong = False
    if strong:
        strong_fail.append(h)
    else:
        weak_fail.append(h)

print("\nSTRONG_FAIL", len(strong_fail))
byf = defaultdict(list)
for h in strong_fail:
    byf[h["file"]].append(h)
for f, fl in sorted(byf.items(), key=lambda x: -len(x[1])):
    print(f"=== {f} ({len(fl)})")
    for it in fl[:15]:
        print(f"  L{it['line']}: {it['text'][:220]}")

print("\nWEAK_FAIL by file (top 40)")
byf2 = defaultdict(list)
for h in weak_fail:
    byf2[h["file"]].append(h)
for f, fl in sorted(byf2.items(), key=lambda x: -len(x[1]))[:40]:
    print(f"--- {f} ({len(fl)})")
    for it in fl[:4]:
        print(f"  L{it['line']}: {it['text'][:200]}")

print("\nCODE-DEBT live")
for h in live:
    if h["class"] == "CODE-DEBT":
        print(f"  {h['file']}:{h['line']}: {h['text'][:180]}")

# Find Complete Director Edition files
print("\n=== Finding Complete Director Edition ===")
for p in ROOT.rglob("*"):
    if p.is_file() and "Complete_Director" in p.name or (
        p.is_file() and "Director" in p.name and "Crucible" in p.name and p.suffix.lower() in (".docx", ".md")
    ):
        print(p.relative_to(ROOT))

# Direct docx scan of key literature for wrong clocks
print("\n=== Targeted docx/md wrong-clock search ===")
try:
    from docx import Document
except ImportError:
    Document = None


def text_of(path: Path) -> str:
    if path.suffix.lower() == ".docx":
        if Document is None:
            return ""
        doc = Document(str(path))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for t in doc.tables:
            for row in t.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        parts.append(cell.text)
        return "\n".join(parts)
    return path.read_text(encoding="utf-8", errors="replace")


targets = []
for pat in [
    "**/The_Directors_Crucible*",
    "**/Chronicle*",
    "**/Instruments*",
    "**/Codex*",
    "**/Matrix*",
    "**/MATRIX*",
    "**/*player rulebook*",
    "**/rulebook_output.txt",
    "**/DESCENT_WINDOW*",
    "**/WORLD_DUNGEON*",
    "**/SESSION_LOOP*",
    "**/NEWCOMER*",
]:
    targets.extend(ROOT.glob(pat))
    targets.extend(ROOT.rglob(pat.replace("**/", "")))

# unique
seen = set()
uniq = []
for t in targets:
    if not t.is_file():
        continue
    s = str(t)
    if s in seen:
        continue
    if "backups" in s.replace("\\", "/"):
        continue
    if re.search(r"\.pre-|pre-hybridb", t.name, re.I):
        continue
    seen.add(s)
    uniq.append(t)

WRONG = re.compile(
    r"(?:240[\-\s]?hour(?:s)?\s+Opening|Opening.{0,50}240\s*hours?|"
    r"ten[\-\s]?day\s+Opening|10[\-\s]?day\s+Opening|"
    r"Horn.{0,40}(?:hour\s*)?120\b(?!0)|hour\s*120(?!\d).{0,40}Horn|"
    r"do not multiply the tabletop clock by ten|"
    r"Director canon.{0,80}240-hour|"
    r"seal closes at hour 240|"
    r"Opening lasts 240|"
    r"measured as 240 hours|"
    r"The 240-Hour Descent|"
    r"have 240 hours|"
    r"240-hour track|"
    r"full 240-hour Opening|"
    r"At hour 120, day 5)",
    re.I,
)
RIGHT = re.compile(r"2400|1200|100[\-\s]?day|day\s*50", re.I)

for path in sorted(uniq, key=lambda p: str(p)):
    if path.suffix.lower() not in (".md", ".txt", ".docx"):
        continue
    try:
        text = text_of(path)
    except Exception as e:
        print(f"ERR {path.relative_to(ROOT)}: {e}")
        continue
    wrong_hits = []
    for i, line in enumerate(text.splitlines(), 1):
        if WRONG.search(line):
            # skip if typo callout
            low = line.lower()
            if any(x in low for x in ("typo", "superseded", "was 240", "drift", "missing zero", "amended", "code debt", "wrongly")):
                continue
            if "2400" in line and "240" in line and re.search(r"240\s*[→\->].*2400|2400.*(was|from).*240", line):
                continue
            wrong_hits.append((i, line.strip()[:240]))
    right_n = len(RIGHT.findall(text))
    rel = str(path.relative_to(ROOT)).replace("\\", "/")
    if wrong_hits or right_n:
        status = "FAIL" if wrong_hits else "PASS-ish"
        print(f"\n[{status}] {rel}  wrong={len(wrong_hits)} right_tokens~{right_n}")
        for i, l in wrong_hits[:20]:
            print(f"  L{i}: {l}")
