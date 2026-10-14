"""T1–5 rules package: XP milestone Opening, DM-01 bible patches, Burnout/ceiling/RP draft labels.

Backs up Player Rulebook + Complete Director Edition before surgical edits.
No commit.
"""
from __future__ import annotations

import shutil
import sys
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parents[2]
PLAYER = ROOT / "decennium descent player rulebook.docx"
DIRECTOR = ROOT / "docs" / "manuals" / "The_Directors_Crucible_Complete_Director_Edition.docx"
BACKUP_DIR = ROOT / "backups"
EXTRACT = ROOT / "rulebook_output.txt"


def set_text(paragraph: Paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for r in paragraph.runs[1:]:
            r.text = ""
    else:
        paragraph.add_run(text)


def insert_after(paragraph: Paragraph, text: str) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    new_para.add_run(text)
    return new_para


def find_prefix(doc: Document, prefix: str) -> Paragraph:
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if t.startswith(prefix):
            return p
    raise SystemExit(f"MISSING prefix: {prefix!r}")


def find_exact(doc: Document, exact: str) -> Paragraph:
    for p in doc.paragraphs:
        if (p.text or "").strip() == exact:
            return p
    raise SystemExit(f"MISSING exact: {exact!r}")


def find_contains(doc: Document, needle: str) -> Paragraph:
    for p in doc.paragraphs:
        if needle in ((p.text or "").strip()):
            return p
    raise SystemExit(f"MISSING contains: {needle!r}")


def backup(path: Path, stamp: str, label: str) -> Path:
    BACKUP_DIR.mkdir(exist_ok=True)
    dest = BACKUP_DIR / f"{label}_{stamp}{path.suffix}"
    shutil.copy2(path, dest)
    print("BACKUP", dest)
    return dest


def patch_player(stamp: str) -> list[str]:
    backup(PLAYER, stamp, "decennium_descent_player_rulebook_pre_t1_5_rules")
    if EXTRACT.exists():
        shutil.copy2(EXTRACT, BACKUP_DIR / f"rulebook_output_pre_t1_5_rules_{stamp}.txt")

    doc = Document(str(PLAYER))
    edits: list[str] = []

    # --- Advancement: Opening = milestone official ---
    table_note = find_prefix(
        doc,
        "The Character Advancement table summarizes the XP you need to advance",
    )
    # Insert after the XP table block — find last table row text then Chapter 2
    ch2 = find_exact(doc, "Chapter 2: Races")
    # Walk backward from ch2 to find last XP row paragraph to insert before ch2
    paras = list(doc.paragraphs)
    i_ch2 = next(i for i, p in enumerate(paras) if (p.text or "").strip() == "Chapter 2: Races")
    anchor = paras[i_ch2 - 1]
    while i_ch2 > 1 and not (paras[i_ch2 - 1].text or "").strip():
        i_ch2 -= 1
        anchor = paras[i_ch2 - 1]

    blocks = [
        "### Advancement in Opening Campaigns",
        (
            "For Opening campaigns (Tier 1–5 play through a Decennial Opening), "
            "milestone advancement is the official method: the Director awards a level "
            "when the Expedition hits a Gate seal, completes a major arc beat, or reaches "
            "another announced milestone. Tell the table which milestones grant levels "
            "before the Opening begins."
        ),
        (
            "The Character Advancement table remains canon for total character level, "
            "proficiency bonus, and for tables that opt into XP tracking. This book does "
            "not publish creature or encounter XP awards. Any companion Session CR→XP "
            "helper is optional and non-canon — it never replaces Director milestones "
            "for Opening campaigns and does not write XP to your character sheet automatically."
        ),
    ]
    cursor = anchor
    for text in blocks:
        cursor = insert_after(cursor, text)
        edits.append(f"player+ {text[:60]}")

    # Clarify Beyond 1st Level lead-in slightly without deleting XP language
    beyond = find_prefix(doc, "As your character goes on adventures and overcomes challenges")
    set_text(
        beyond,
        (
            "As your character goes on adventures and overcomes challenges, they gain "
            "experience. Opening campaigns use milestone advancement (see Advancement "
            "in Opening Campaigns below). Experience points and the Character Advancement "
            "table remain available when your Director opts into XP tracking: a character "
            "who reaches a specified experience point total advances in capability. This "
            "advancement is called gaining a level."
        ),
    )
    edits.append("player~ Beyond 1st Level lead-in")

    # --- Non-core RP pools draft ---
    rp_note = find_prefix(doc, "*Note: These base values are modified by primary spellcasting")
    set_text(
        rp_note,
        (
            "*Note: These base values are modified by primary spellcasting attribute, "
            "alignment affinity, and reputation bonuses. Wizard, Mage, Cleric, and Druid "
            "tables above are rulebook-complete. Other Weave-casting classes use draft "
            "companion RP pool estimates (clearly labeled Draft in the companion) until "
            "author tables are published — treat those numbers as playtest material, not "
            "final law.*"
        ),
    )
    edits.append("player~ RP pool draft note")

    # --- Burnout GM-fiat ---
    burnout = find_prefix(doc, "Overcasting risks Resonance Burnout")
    set_text(
        burnout,
        (
            "Overcasting risks Resonance Burnout if a caster depletes their pool below a "
            "critical threshold (25% of maximum). Until a final Burnout consequences table "
            "is published, Burnout effects are Director fiat: announce the consequence "
            "before it applies (the Director's Crucible suggests disadvantage on "
            "Weave-related checks until the end of the caster's next turn as a default "
            "playtest ruling). Companion Session Burnout badges are optional reminders, "
            "not automated rules."
        ),
    )
    edits.append("player~ Burnout GM-fiat")

    # --- Class ceiling caps note (after Overcasting / before Alignment) ---
    align = find_exact(doc, "## Alignment and Magic")
    # insert before Alignment — find Overcasting burnout para again after edits
    burnout2 = find_prefix(doc, "Overcasting risks Resonance Burnout")
    insert_after(
        burnout2,
        (
            "Class ceiling caps (action-economy limits for high-impact features such as "
            "extra mutations, Revelations, or smite riders) are draft playtest guardrails "
            "published for Directors. They are GM fiat until promoted: if your table uses "
            "them, the Director announces the relevant cap; companion callouts are "
            "informational only and do not auto-enforce."
        ),
    )
    edits.append("player+ ceiling caps GM-fiat")

    doc.save(str(PLAYER))
    print("SAVED", PLAYER)

    # Refresh UTF-16LE extract
    lines = []
    for p in Document(str(PLAYER)).paragraphs:
        lines.append(p.text or "")
    EXTRACT.write_text("\n".join(lines) + "\n", encoding="utf-16-le")
    print("REFRESHED", EXTRACT)
    return edits


def patch_director(stamp: str) -> list[str]:
    backup(DIRECTOR, stamp, "Complete_Director_Edition_pre_t1_5_rules")
    doc = Document(str(DIRECTOR))
    edits: list[str] = []

    # Treasure / XP awards → Opening milestone official
    awards = find_prefix(doc, "Treasure and awards. The cumulative level 1–20 XP thresholds")
    set_text(
        awards,
        (
            "Treasure and awards. The cumulative level 1–20 XP thresholds in the Player "
            "Rulebook are canon for level and proficiency when XP tracking is in force. "
            "For Opening campaigns (Tier 1–5 through a Decennial Opening), milestone "
            "advancement is official: award levels at Gate seals, major arc beats, or "
            "other milestones announced before the Opening begins. No published CR→XP "
            "or encounter-budget table exists in this book or the Player Rulebook. "
            "Companion Session CR→XP helpers are optional and non-canon — label them "
            "playtest if used; they do not write to character saves. Outside Opening "
            "scope, specific encounter awards remain Director judgment unless a published "
            "adventure provides them."
        ),
    )
    edits.append("director~ Treasure and awards")

    # Burnout — permanent GM-fiat label
    burnout = find_prefix(doc, "Overcast and Burnout.")
    set_text(
        burnout,
        (
            "Overcast and Burnout. Extended, Subtle, and Distant each cost +1 RP; "
            "Empowered costs +2; Widened, Quickened, and Twinned cost +3; Heightened "
            "costs +4. Burnout risk begins when the remaining pool falls below one "
            "quarter of maximum after a cast. Until a final Burnout table is published, "
            "Burnout is permanent GM fiat: announce a playtest consequence before it "
            "applies (default suggestion: disadvantage on Weave-related checks until "
            "the end of the caster's next turn) and record how it affected play. "
            "Companion Session Burnout badges are optional reminders only."
        ),
    )
    edits.append("director~ Burnout GM-fiat")

    # RP pool draft clarity
    pool = find_prefix(doc, "Pool size. Use the class tables in the Player Rulebook.")
    set_text(
        pool,
        (
            "Pool size. Use the class tables in the Player Rulebook. Wizard, Mage, "
            "Cleric, and Druid have complete level 1–20 Resonance tables. When a "
            "Weave-casting class lacks a finalized table, treat companion draft RP "
            "pool estimates as playtest material only — write the chosen progression "
            "on the character sheet and say Draft aloud at session zero."
        ),
    )
    edits.append("director~ RP pool draft")

    # Ch 21 — insert ceiling / burnout adjudication block after "When a ruling interacts..."
    ruling = find_prefix(doc, "When a ruling interacts with a class limit.")
    set_text(
        ruling,
        (
            "When a ruling interacts with a class limit. Let the most specific published "
            "class feature, action-economy limit, copy limit, or aura limit control for "
            "that session. Draft class ceiling caps (companion Class Ceiling Caps callouts "
            "and CLASS_CEILING_CAPS.md) are permanent GM-fiat playtest guardrails until "
            "author-promoted — announce them before they bind a player, record ambiguity, "
            "and clarify before the next session instead of relitigating mid-fight."
        ),
    )
    edits.append("director~ ceiling caps GM-fiat")

    # Silent Ones = Old Ones explicit (Ch 8)
    silent = find_contains(
        doc,
        "The Silent Ones — the Old Ones called the Ul'nak after the Sundering",
    )
    old = (silent.text or "").strip()
    if "Silent Ones ARE the Old Ones" not in old:
        set_text(
            silent,
            (
                "🔒 Director-secret. Silent Ones ARE the Old Ones (same entity). "
                "Register note: Silent Ones is the GM-secret framing; Old Ones / Ul'nak "
                "is the Matrix-of-Oblivion cosmology framing — not a second force. "
                "The Silent Ones — the Old Ones called the Ul'nak after the Sundering — "
                "built the machine beneath the public Crucible: the Heartcell. It contains "
                "the hundred-floor resonance structure, the ten Gate divisions, the "
                "recurring Opening, and the reward architecture that makes loot behave as "
                "bait. Calan'thir and Asmodeus did not originate that machine. After the "
                "Old Gods were barely defeated, they found it, claimed it, modified it, "
                "activated it as a resonance valve, and bound the Divine Pact to its "
                "operation. Both histories are therefore true at different layers: the "
                "Silent Ones built the prison; Calan'thir and Asmodeus built the Crucible "
                "the world enters."
            ),
        )
        edits.append("director~ Silent Ones = Old Ones (Ch 8)")

    # Canon Guardrails
    guard = find_prefix(
        doc,
        "Director-secret: the Silent Ones, later called the Ul",
    )
    set_text(
        guard,
        (
            "Director-secret: Silent Ones ARE the Old Ones (same entity). The Silent "
            "Ones, later called the Ul'nak, built the underlying Heartcell."
        ),
    )
    edits.append("director~ Canon Guardrails Silent Ones")

    # Glossary
    gloss = find_exact(doc, "Silent Ones / Ul’nak")
    # next meaningful para is the definition — find by scanning
    found = False
    paras = list(doc.paragraphs)
    for i, p in enumerate(paras):
        t = (p.text or "").strip()
        if t in ("Silent Ones / Ul’nak", "Silent Ones / Ul'nak"):
            # definition is usually next non-empty
            for j in range(i + 1, min(i + 4, len(paras))):
                d = (paras[j].text or "").strip()
                if d and not d.startswith("Matrix"):
                    set_text(
                        paras[j],
                        (
                            "Director-secret builders; Silent Ones ARE the Old Ones "
                            "(same entity). Ul'nak is the post-Sundering / Matrix framing "
                            "of that same force."
                        ),
                    )
                    edits.append("director~ glossary Silent Ones")
                    found = True
                    break
            break
    if not found:
        print("WARN: glossary Silent Ones definition not updated")

    # Always checklist — XP awards line already says provisional; tighten
    always_xp = find_prefix(doc, "Label provisional combat statistics, XP awards")
    set_text(
        always_xp,
        (
            "Label provisional combat statistics and any optional CR→XP companion helpers "
            "as playtest / non-canon. Opening campaigns use milestone advancement as "
            "official; do not present Session draft XP as rulebook awards."
        ),
    )
    edits.append("director~ Always checklist XP")

    doc.save(str(DIRECTOR))
    print("SAVED", DIRECTOR)
    return edits


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8")
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    pe = patch_player(stamp)
    de = patch_director(stamp)
    print("PLAYER EDITS", len(pe))
    print("DIRECTOR EDITS", len(de))


if __name__ == "__main__":
    main()
