#!/usr/bin/env python3
"""Targeted wrong-clock scan of key live literature."""
from __future__ import annotations

import re
import sys
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

ROOT = Path(r"C:\Github\Descent assets\DecenniumDescent-main")
from docx import Document

WRONG = re.compile(
    r"(?:240[\-\s]?hour(?:s)?\s+[Oo]pening|[Oo]pening.{0,60}240\s*hours?\b|"
    r"ten[\-\s]?day\s+[Oo]pening|10[\-\s]?day\s+[Oo]pening|"
    r"[Hh]orn.{0,50}(?:hour\s*)?120\b(?!0)|(?:hour\s*)?120(?!\d).{0,50}[Hh]orn|"
    r"do not multiply the tabletop clock by ten|"
    r"Director canon.{0,100}240-hour|"
    r"seal closes at hour 240|"
    r"[Oo]pening lasts 240|"
    r"measured as 240 hours|"
    r"The 240-Hour Descent|"
    r"have 240 hours|"
    r"240-hour track|"
    r"full 240-hour Opening|"
    r"At hour 120, day 5|"
    r"Horn sounds at hour 120\b|"
    r"blank 240-hour|"
    r"Current Opening hour and a blank 240)",
    re.I,
)

# More precise for live canon assertions of wrong Opening (not horn code debt)
WRONG_OPENING = re.compile(
    r"(?:240[\-\s]?hour(?:s)?\s+[Oo]pening|[Oo]pening.{0,40}(?:is|lasts|of)\s*240\s*hours?|"
    r"ten[\-\s]?day\s+[Oo]pening|10[\-\s]?day\s+[Oo]pening|"
    r"do not multiply the tabletop clock by ten|"
    r"Director canon for the tabletop game is a 240|"
    r"seal closes at hour 240|"
    r"The 240-Hour Descent|"
    r"have 240 hours|"
    r"240-hour track|"
    r"full 240-hour Opening|"
    r"measured as 240 hours|"
    r"blank 240-hour)",
    re.I,
)

WRONG_HORN_LIVE = re.compile(
    r"(?:At hour 120, day 5|"
    r"Horn sounds at hour 120\b|"
    r"Horn at hour 120\b|"
    r"with the Horn at hour 120\b|"
    r"Horn PoNR.{0,20}120\b(?!0))",
    re.I,
)

CORRECT = re.compile(
    r"2400|1200\s*h|hour\s*1200|1200\s*hours|100[\-\s]?days?|day\s*50",
    re.I,
)

TYPO_OK = re.compile(
    r"typo|superseded|drift|wrongly|missing zero|amended|code debt|code/doc debt|"
    r"was\s*240|still reads|do not change|pending|ask owner|historical|"
    r"240\s*[→\-–>]+\s*2400|2400.{0,30}(?:was|from).{0,10}240|"
    r"may be wrong|flag(?:ged)?|not changed|intentionally left",
    re.I,
)


def read_docx(path: Path) -> str:
    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)
    return "\n".join(parts)


def read_any(path: Path) -> str:
    if path.suffix.lower() == ".docx":
        return read_docx(path)
    raw = path.read_bytes()
    for enc in ("utf-8", "utf-16", "utf-16-le", "latin-1"):
        try:
            return raw.decode(enc)
        except Exception:
            continue
    return ""


PATHS = [
    ROOT / "docs/manuals/The_Directors_Crucible_Complete_Director_Edition.docx",
    ROOT / "docs/manuals/Chronicle_of_the_Decennial_Opening.docx",
    ROOT / "docs/manuals/Chronicle_of_the_Decennial_Opening.md",
    ROOT / "docs/manuals/Instruments_of_the_Weave.md",
    ROOT / "docs/manuals/Instruments_of_the_Weave.docx",
    ROOT / "docs/manuals/Codex_of_the_Heartcell.md",
    ROOT / "docs/manuals/Codex_of_the_Heartcell.docx",
    ROOT / "docs/manuals/Decennium_Descent_Directors_Manual.md",
    ROOT / "docs/manuals/Decennium_Descent_Directors_Manual.docx",
    ROOT / "docs/manuals/DIRECTOR_EDITION_BIBLE_REVIEW.md",
    ROOT / "docs/manuals/AGENT_SIMULATION_REFERENCE.md",
    ROOT / "The Decennium Descent and the Matrix of Oblivion.docx",
    ROOT / "docs/playtest/MATRIX_OF_OBLIVION_EXTRACT.md",
    ROOT / "decennium descent player rulebook.docx",
    ROOT / "rulebook_output.txt",
    ROOT / "docs/mechanics/DESCENT_WINDOW_CANON.md",
    ROOT / "docs/mechanics/WORLD_DUNGEON_LOOP.md",
    ROOT / "docs/mechanics/SESSION_LOOP.md",
    ROOT / "docs/mechanics/NEWCOMER_GM_GUIDE.md",
    ROOT / "docs/MECHANICS_BIBLE_STUB.md",
    ROOT / "README.md",
    ROOT / "src/data/dungeon-mechanics.json",
    ROOT / "src/data/floor-encounters.json",
    ROOT / "scripts/floor-encounters.json",
]

for path in PATHS:
    if not path.exists():
        print(f"MISSING {path.relative_to(ROOT)}")
        continue
    text = read_any(path)
    lines = text.splitlines()
    open_fails = []
    horn_fails = []
    ok_ctx = []
    pass_lines = []
    for i, line in enumerate(lines, 1):
        if CORRECT.search(line) and re.search(r"opening|horn|2400|1200|ponr", line, re.I):
            if not WRONG_OPENING.search(line) and not WRONG_HORN_LIVE.search(line):
                pass_lines.append(i)
        if WRONG_OPENING.search(line) or WRONG_HORN_LIVE.search(line) or (
            WRONG.search(line) and re.search(r"opening|240-hour|ten-day|multiply", line, re.I)
        ):
            if TYPO_OK.search(line):
                ok_ctx.append((i, line.strip()[:220]))
            elif WRONG_OPENING.search(line) or WRONG_HORN_LIVE.search(line):
                open_fails.append((i, line.strip()[:220]))
            elif re.search(r"\b120\s*h|hour\s*120\b|HORN_AT\s*=\s*120|120 hours underground", line, re.I):
                horn_fails.append((i, line.strip()[:220]))

    # also catch hornRule product strings without Opening
    for i, line in enumerate(lines, 1):
        if re.search(r"At 120 hours(?: underground)?, a giant horn", line, re.I):
            if not any(x[0] == i for x in open_fails + horn_fails + ok_ctx):
                horn_fails.append((i, line.strip()[:220]))

    rel = str(path.relative_to(ROOT)).replace("\\", "/")
    print(f"\n==== {rel}")
    print(f"  PASS-ish clock lines: {len(pass_lines)}  OPENING/HORN FAIL: {len(open_fails)}  HORN-120 claims: {len(horn_fails)}  OK-CONTEXT: {len(ok_ctx)}")
    for i, l in open_fails[:25]:
        print(f"  FAIL L{i}: {l}")
    for i, l in horn_fails[:15]:
        print(f"  HORN120 L{i}: {l}")
    for i, l in ok_ctx[:8]:
        print(f"  OKCTX L{i}: {l}")
    # sample correct
    for i in pass_lines[:3]:
        print(f"  PASS L{i}: {lines[i-1].strip()[:180]}")
